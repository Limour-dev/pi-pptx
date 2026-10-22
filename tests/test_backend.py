#!/usr/bin/env python3
"""Backend smoke tests for python/pptkit.py.

Run: python3 tests/test_backend.py
Set PI_PPTX_TEST_TEMPLATE to also test against a real deck (optional).
"""
from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
KIT = ROOT / "python" / "pptkit.py"

# 1x1 transparent PNG
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def kit(*args: str, expect_ok: bool = True) -> dict:
    result = subprocess.run(
        [sys.executable, str(KIT), *args],
        capture_output=True,
        text=True,
        timeout=300,
    )
    if expect_ok:
        if result.returncode != 0:
            raise AssertionError(f"pptkit {' '.join(args)} failed: {result.stderr}")
        return json.loads(result.stdout)
    if result.returncode == 0:
        raise AssertionError(f"pptkit {' '.join(args)} was expected to fail")
    return {"stderr": result.stderr}


class PptkitTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="pi-pptx-test-"))
        self.deck = self.tmp / "deck.pptx"
        self.png = self.tmp / "pixel.png"
        self.png.write_bytes(PNG_1PX)

    def create(self, name: str = "deck.pptx", **kw) -> Path:
        deck = self.tmp / name
        args = ["create", str(deck)]
        if kw.get("template"):
            args += ["--template", kw["template"]]
        kit(*args)
        return deck

    def inspect(self, deck: Path, slides: str | None = None, **kw) -> dict:
        args = ["inspect", str(deck)]
        if slides:
            args += ["--slides", slides]
        for key, value in kw.items():
            args += ["--" + key.replace("_", "-"), str(value)]
        return kit(*args)

    def make_template(self) -> Path:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Sample title"
        slide.shapes.add_picture(str(self.png), Inches(1), Inches(1), Inches(2), Inches(2))
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3.5), Inches(4), Inches(1.2)).table
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "EF"
        table.cell(1, 1).text = "55"
        template = self.tmp / "template.pptx"
        prs.save(template)
        return template

    # ------------------------------------------------------------------ tests
    def test_doctor(self):
        info = kit("doctor")
        self.assertIn("python_pptx", info)
        self.assertNotIn("ERROR", str(info.get("python_pptx")))
        self.assertIn("pymupdf", info)

    def test_create_blank_and_overwrite_guard(self):
        deck = self.create()
        info = self.inspect(deck)
        self.assertEqual(info["slide_count"], 1)
        self.assertEqual(info["slide_size_in"], [13.333, 7.5])
        failure = kit("create", str(deck), expect_ok=False)
        self.assertIn("已存在", failure["stderr"])
        kit("create", str(deck), "--overwrite")
        self.assertEqual(self.inspect(deck)["slide_count"], 1)

    def test_create_from_template_copies_slides(self):
        template = self.make_template()
        deck = self.create("copy.pptx", template=template)
        info = self.inspect(deck)
        self.assertEqual(info["slide_count"], 1)
        types = [s["type"] for s in info["slides"][0]["shapes"]]
        self.assertIn("PICTURE", types)
        self.assertIn("TABLE", types)

    def test_clone_preserves_pictures_table_and_notes(self):
        from pptx import Presentation

        template = self.make_template()
        deck = self.create("copy.pptx", template=template)
        # add notes to the source page
        prs = Presentation(str(deck))
        prs.slides[0].notes_slide.notes_text_frame.text = "speaker note"
        prs.save(str(deck))

        out = kit("clone", str(deck), "1", "--at", "1", "--count", "1")
        self.assertEqual(out["slide_count"], 2)
        info = self.inspect(deck)
        first = info["slides"][0]["shapes"]
        self.assertIn("PICTURE", [s["type"] for s in first])
        table = next(s for s in first if s.get("table"))
        self.assertEqual(table["table"][1][1], "55")
        self.assertEqual(info["slides"][0].get("notes"), "speaker note")

    def test_set_text_preserves_style_and_paragraphs(self):
        deck = self.create()
        kit(
            "add-text",
            str(deck),
            "--slide",
            "1",
            "--text",
            "原始标题",
            "--x",
            "1",
            "--y",
            "1",
            "--w",
            "8",
            "--h",
            "1",
            "--size",
            "28",
            "--font",
            "微软雅黑",
            "--color",
            "#1F4E79",
            "--bold",
        )
        info = self.inspect(deck, "1")
        shape = info["slides"][0]["shapes"][-1]
        name = shape["name"]
        self.assertEqual(shape["style"]["size_pt"], 28.0)
        self.assertEqual(shape["style"]["font"], "微软雅黑")

        kit("set-text", str(deck), "--slide", "1", "--shape", name, "--text", "新标题\n第二段")
        info = self.inspect(deck, "1")
        shape = info["slides"][0]["shapes"][-1]
        self.assertEqual(shape["text"], "新标题\n第二段")
        self.assertEqual(shape["style"]["size_pt"], 28.0)
        self.assertEqual(shape["style"]["font"], "微软雅黑")

    def test_table_add_and_set_cell(self):
        deck = self.create()
        added = kit(
            "add-table",
            str(deck),
            "--slide",
            "1",
            "--rows",
            "2",
            "--cols",
            "2",
            "--x",
            "1",
            "--y",
            "1",
            "--w",
            "4",
            "--h",
            "1.5",
            "--data",
            '[["指标","值"],["EF","52"]]',
        )
        name = added["name"]
        kit("set-cell", str(deck), "--slide", "1", "--shape", name, "--row", "1", "--col", "1", "--text", "58")
        info = self.inspect(deck, "1")
        table = next(s for s in info["slides"][0]["shapes"] if s.get("table"))
        self.assertEqual(table["table"], [["指标", "值"], ["EF", "58"]])
        kit("set-cell", str(deck), "--slide", "1", "--shape", name, "--row", "9", "--col", "0", "--text", "x", expect_ok=False)

    def test_add_image_and_group_path_error(self):
        deck = self.create()
        added = kit(
            "add-image",
            str(deck),
            "--slide",
            "1",
            "--path",
            str(self.png),
            "--x",
            "2",
            "--y",
            "2",
            "--w",
            "1",
        )
        self.assertEqual(added["added"], "picture")
        self.assertAlmostEqual(added["box_in"][2], 1.0, places=1)
        failure = kit(
            "set-text",
            str(deck),
            "--slide",
            "1",
            "--shape",
            "不存在的形状",
            "--text",
            "x",
            expect_ok=False,
        )
        self.assertIn("可用形状路径", failure["stderr"])

    def test_shot_render_overflow_and_cache(self):
        deck = self.create()
        kit("add-text", str(deck), "--slide", "1", "--text", "正常的一行", "--x", "1", "--y", "1", "--w", "4", "--h", "1", "--size", "20")
        bad = kit(
            "add-text",
            str(deck),
            "--slide",
            "1",
            "--text",
            "这是一段很长很长的文字，用来制造确定性的溢出，" * 3,
            "--x",
            "1",
            "--y",
            "2.5",
            "--w",
            "4",
            "--h",
            "0.3",
            "--size",
            "18",
        )
        first = kit(
            "shot",
            str(deck),
            "--slides",
            "1",
            "--dpi",
            "90",
        )
        self.assertEqual(len(first["shots"]), 1)
        shot = first["shots"][0]
        self.assertTrue(Path(shot["file"]).exists())
        self.assertGreater(shot["px"][0], 100)
        names = [item["shape"] for item in shot.get("overflow", [])]
        self.assertIn(bad["name"], names)
        texts = [item["text"] for item in shot.get("overflow", [])]
        self.assertFalse(any("正常的一行" in text for text in texts))

        # second shot hits the PDF cache but still works
        second = kit("shot", str(deck), "--slides", "1", "--dpi", "90")
        self.assertEqual(second["shots"][0]["file"], shot["file"])
        self.assertEqual(second["pdf"], first["pdf"])

    def test_shot_max_pages_and_delete(self):
        deck = self.create()
        for _ in range(3):
            kit("clone", str(deck), "1")
        result = kit("shot", str(deck), "--dpi", "80", "--max-pages", "2")
        self.assertEqual([s["slide"] for s in result["shots"]], [1, 2])
        self.assertEqual(result["omitted_slides"], [3, 4])
        deleted = kit("delete", str(deck), "--slides", "2-3")
        self.assertEqual(deleted["slide_count"], 2)
        kit("delete", str(deck), "--slides", "1-2", expect_ok=False)

    def test_add_slide_by_name(self):
        deck = self.create()
        info = self.inspect(deck)
        layout = info["layouts"][0]
        out = kit("add-slide", str(deck), "--layout", layout["name"])
        self.assertEqual(out["layout"], layout["name"])
        self.assertEqual(out["added_slide"], 2)

    def slide_xml(self, deck: Path, slide: int = 1) -> str:
        with zipfile.ZipFile(deck) as archive:
            return archive.read(f"ppt/slides/slide{slide}.xml").decode("utf-8")

    def test_add_text_cjk_font_layout_and_autofit(self):
        deck = self.create()
        added = kit(
            "add-text",
            str(deck),
            "--slide",
            "1",
            "--text",
            "中文探测行\n第二段",
            "--x",
            "1",
            "--y",
            "1",
            "--w",
            "6",
            "--size",
            "20",
            "--font",
            "微软雅黑",
            "--space-after",
            "6",
            "--indent",
            "0.32",
            "--bullet",
            "•",
            "--bullet-color",
            "#2E75B6",
        )
        # h omitted -> grow + estimated height
        self.assertEqual(added["fit"], "grow")
        self.assertGreater(added["box_in"][3], 0.8)
        xml = self.slide_xml(deck)
        self.assertIn('<a:ea typeface="微软雅黑"/>', xml)
        self.assertIn('<a:cs typeface="微软雅黑"/>', xml)
        self.assertIn("<a:spAutoFit/>", xml)
        self.assertIn("<a:buChar", xml)
        self.assertIn("<a:buClr>", xml)
        self.assertIn('marL="292608"', xml)
        self.assertIn('indent="-292608"', xml)
        self.assertIn('<a:spcAft>', xml)

        # explicit h + fit=shrink writes normAutofit, no grow
        shrunk = kit(
            "add-text",
            str(deck),
            "--slide",
            "1",
            "--text",
            "固定框",
            "--x",
            "1",
            "--y",
            "3",
            "--w",
            "3",
            "--h",
            "1",
            "--fit",
            "shrink",
        )
        self.assertEqual(shrunk["fit"], "shrink")
        self.assertIn("<a:normAutofit", self.slide_xml(deck))

    def test_set_text_rich_runs_keep_base_style(self):
        deck = self.create()
        added = kit(
            "add-text",
            str(deck),
            "--slide",
            "1",
            "--text",
            "原始",
            "--x",
            "1",
            "--y",
            "1",
            "--w",
            "8",
            "--h",
            "1.5",
            "--size",
            "20",
            "--font",
            "微软雅黑",
            "--color",
            "#1F1F1F",
        )
        paras = json.dumps(
            [
                {
                    "runs": [
                        {"t": "一句实话：", "bold": True, "color": "C00000", "size": 22},
                        {"t": "没有事前效能计算。"},
                    ],
                    "space_after": 6,
                },
                {"runs": [{"t": "第二段"}], "align": "center"},
            ],
            ensure_ascii=False,
        )
        out = kit("set-text", str(deck), "--slide", "1", "--shape", "#" + str(added["id"]), "--paras", paras)
        self.assertEqual(out["paragraphs"], 2)

        from pptx import Presentation

        prs = Presentation(str(deck))
        shape = prs.slides[0].shapes[-1]
        first, second = shape.text_frame.paragraphs
        bold, normal = first.runs
        self.assertEqual(bold.text, "一句实话：")
        self.assertTrue(bold.font.bold)
        self.assertEqual(bold.font.size.pt, 22.0)
        self.assertEqual(str(bold.font.color.rgb), "C00000")
        # base run style inherited where not overridden
        self.assertEqual(normal.text, "没有事前效能计算。")
        self.assertEqual(normal.font.size.pt, 20.0)
        self.assertEqual(normal.font.name, "微软雅黑")
        self.assertAlmostEqual(first.space_after.pt, 6.0)
        self.assertEqual(str(second.alignment), "CENTER (2)")

    def test_shape_id_addressing_duplicate_names_and_delete(self):
        deck = self.create()
        first = kit("add-text", str(deck), "--slide", "1", "--text", "A", "--x", "1", "--y", "1", "--w", "3", "--h", "1")
        second = kit("add-text", str(deck), "--slide", "1", "--text", "B", "--x", "5", "--y", "1", "--w", "3", "--h", "1")

        from pptx import Presentation

        prs = Presentation(str(deck))
        for shape in list(prs.slides[0].shapes)[-2:]:
            shape.name = "标题 3"
        prs.save(str(deck))

        failure = kit(
            "set-text", str(deck), "--slide", "1", "--shape", "标题 3", "--text", "x", expect_ok=False
        )
        self.assertIn("不唯一", failure["stderr"])
        self.assertIn(f"id=#{first['id']}", failure["stderr"])
        self.assertIn(f"id=#{second['id']}", failure["stderr"])
        self.assertIn("index=", failure["stderr"])

        out = kit("set-text", str(deck), "--slide", "1", "--shape", f"#{second['id']}", "--text", "B2")
        self.assertEqual(out["text"], "B2")
        info = self.inspect(deck, "1")
        texts = {shape["id"]: shape.get("text") for shape in info["slides"][0]["shapes"]}
        self.assertEqual(texts[first["id"]], "A")
        self.assertEqual(texts[second["id"]], "B2")

        # out-of-range numeric index explains the 0-based convention
        failure = kit(
            "set-text", str(deck), "--slide", "1", "--shape", "99", "--text", "x", expect_ok=False
        )
        self.assertIn("index 从 0 开始", failure["stderr"])

        deleted = kit("delete-shape", str(deck), "--slide", "1", "--shape", f"#{first['id']}")
        self.assertTrue(deleted["deleted"])
        ids = [shape["id"] for shape in self.inspect(deck, "1")["slides"][0]["shapes"]]
        self.assertNotIn(first["id"], ids)

    def test_notes_summary_inspect_and_doctor_theme(self):
        deck = self.create()
        out = kit("set-notes", str(deck), "--slide", "1", "--text", "开场白\n第二段")
        self.assertEqual(out["paragraphs"], 2)
        summary = self.inspect(deck, "1", mode="summary")
        self.assertTrue(summary["slides"][0]["title"])
        self.assertIn("开场白", summary["slides"][0]["notes"])
        self.assertEqual(summary["slides"][0]["shape_count"], 2)

        full = kit("inspect", str(deck), "--slides", "1", "--depth", "0")
        self.assertNotIn("children", json.dumps(full["slides"], ensure_ascii=False))
        doctor = kit("doctor", str(deck))
        self.assertIn("deck_theme_fonts", doctor)
        self.assertIn("major", doctor["deck_theme_fonts"])

    def test_table_theme_off_style_and_borders(self):
        deck = self.create()
        added = kit(
            "add-table",
            str(deck),
            "--slide",
            "1",
            "--rows",
            "2",
            "--cols",
            "2",
            "--x",
            "1",
            "--y",
            "3",
            "--w",
            "5",
            "--h",
            "1.5",
            "--data",
            '[["指标","值"],["EF","52"]]',
            "--size",
            "13",
            "--header-fill",
            "#1F4E79",
            "--header-color",
            "#FFFFFF",
            "--zebra",
            "#EAF1F8",
        )
        self.assertIn("模板里没有任何表格样例", added["hint"])

        from pptx import Presentation
        from pptx.oxml.ns import qn

        prs = Presentation(str(deck))
        shape = next(s for s in prs.slides[0].shapes if s.has_table)
        tbl = shape._element.graphic.graphicData.tbl
        tblPr = tbl.find(qn("a:tblPr"))
        self.assertEqual(tblPr.get("firstRow"), "0")
        self.assertEqual(tblPr.get("bandRow"), "0")
        self.assertIsNone(tblPr.find(qn("a:tableStyleId")))
        header = shape.table.cell(0, 0)
        self.assertEqual(str(header.fill.fore_color.rgb), "1F4E79")
        body = shape.table.cell(1, 0)
        self.assertEqual(str(body.fill.fore_color.rgb), "EAF1F8")
        self.assertEqual([el.tag.split("}")[1] for el in list(body._tc.tcPr)[:4]], ["lnL", "lnR", "lnT", "lnB"])

        kit(
            "table-style",
            str(deck),
            "--slide",
            "1",
            "--shape",
            f"#{added['id']}",
            "--header-fill",
            "#C00000",
            "--border",
            "none",
            "--col-widths",
            "2,3",
        )
        prs = Presentation(str(deck))
        shape = next(s for s in prs.slides[0].shapes if s.has_table)
        self.assertEqual(str(shape.table.cell(0, 1).fill.fore_color.rgb), "C00000")
        self.assertAlmostEqual(shape.table.columns[0].width.inches, 2.0, places=2)
        lnL = shape.table.cell(1, 1)._tc.tcPr.find(qn("a:lnL"))
        self.assertIsNotNone(lnL.find(qn("a:noFill")))

    def test_apply_dry_run_atomic_failure_and_commit(self):
        deck = self.create()
        added = kit("add-text", str(deck), "--slide", "1", "--text", "初始", "--x", "1", "--y", "1", "--w", "5", "--h", "1")
        before = hashlib.sha256(deck.read_bytes()).hexdigest()

        good_ops = [
            {"op": "set_text", "slide": 1, "shape": f"#{added['id']}", "text": "批量改完"},
            {"op": "add_text", "slide": 1, "x": 1, "y": 3, "w": 5, "text": "新增", "size": 16},
            {"op": "set_notes", "slide": 1, "text": "备注"},
        ]
        dry = kit("apply", str(deck), "--ops", json.dumps(good_ops, ensure_ascii=False), "--dry-run")
        self.assertTrue(dry["dry_run"])
        self.assertFalse(dry["saved"])
        self.assertTrue(dry["ok"])
        self.assertEqual(dry["applied"], 3)
        self.assertEqual(hashlib.sha256(deck.read_bytes()).hexdigest(), before)
        self.assertEqual([s.get("text") for s in self.inspect(deck, "1")["slides"][0]["shapes"]], ["标题", "副标题", "初始"])

        bad_ops = [
            {"op": "set_text", "slide": 1, "shape": f"#{added['id']}", "text": "不应写入"},
            {"op": "set_text", "slide": 1, "shape": "#999", "text": "x"},
        ]
        bad = kit("apply", str(deck), "--ops", json.dumps(bad_ops), "--no-shot")
        self.assertFalse(bad["ok"])
        self.assertFalse(bad["saved"])
        self.assertEqual(bad["errors"][0]["index"], 1)
        self.assertEqual(hashlib.sha256(deck.read_bytes()).hexdigest(), before)

        applied = kit("apply", str(deck), "--ops", json.dumps(good_ops, ensure_ascii=False), "--no-shot")
        self.assertTrue(applied["ok"])
        self.assertTrue(applied["saved"])
        self.assertEqual(applied["changed_slides"], [1])
        info = self.inspect(deck, "1")
        texts = [s.get("text") for s in info["slides"][0]["shapes"]]
        self.assertIn("批量改完", texts)
        self.assertIn("新增", texts)
        self.assertEqual(info["slides"][0]["notes"], "备注")

    def test_apply_auto_shot_reports_overflow(self):
        deck = self.create()
        added = kit("add-text", str(deck), "--slide", "1", "--text", "短", "--x", "1", "--y", "1", "--w", "2", "--h", "0.3", "--size", "18")
        ops = [
            {
                "op": "set_text",
                "slide": 1,
                "shape": f"#{added['id']}",
                "text": "这是一段很长很长的文字，" * 4,
            }
        ]
        out = kit("apply", str(deck), "--ops", json.dumps(ops, ensure_ascii=False))
        self.assertTrue(out["saved"])
        self.assertIn("overflow", out)
        names = [item["shape"] for entry in out["overflow"] for item in entry.get("overflow", [])]
        self.assertIn(added["name"], names)

    def test_shot_montage_region_and_overlap(self):
        deck = self.create()
        kit("add-text", str(deck), "--slide", "1", "--text", "重叠一", "--x", "1", "--y", "2", "--w", "4", "--h", "1", "--size", "20")
        kit("add-text", str(deck), "--slide", "1", "--text", "重叠二", "--x", "1.2", "--y", "2.1", "--w", "4", "--h", "1", "--size", "20")
        kit("clone", str(deck), "1", "--count", "2")
        shot = kit(
            "shot",
            str(deck),
            "--slides",
            "1-3",
            "--dpi",
            "70",
            "--montage",
            "2",
            "--overflow",
            "off",
        )
        self.assertEqual(shot["overflow_mode"], "off")
        self.assertEqual(len(shot["shots"]), 3)
        montage = shot["montage"]
        self.assertEqual(montage["cols"], 2)
        self.assertEqual(montage["rows"], 2)
        self.assertTrue(Path(montage["file"]).exists())
        self.assertGreater(montage["px"][0], 600)
        self.assertTrue(shot["shots"][0]["overlaps"])

        region = kit(
            "shot",
            str(deck),
            "--slides",
            "1",
            "--dpi",
            "90",
            "--region",
            "0.5,0.5,3,2",
            "--overflow",
            "off",
        )
        self.assertEqual(region["region_in"], [0.5, 0.5, 3.0, 2.0])
        self.assertLess(region["shots"][0]["px"][0], 400)

    def test_loose_overflow_skips_centered_labels(self):
        deck = self.create()
        added = kit(
            "add-text",
            str(deck),
            "--slide",
            "1",
            "--text",
            "LVEF，%",
            "--x",
            "2",
            "--y",
            "2",
            "--w",
            "2",
            "--h",
            "0.6",
            "--size",
            "18",
            "--align",
            "center",
        )
        shot = kit("shot", str(deck), "--slides", "1", "--dpi", "90", "--overflow", "loose")
        flagged = [item["shape"] for item in shot["shots"][0].get("overflow", [])]
        self.assertNotIn(added["name"], flagged)

    def test_real_template_with_pictures(self):
        template = os.environ.get("PI_PPTX_TEST_TEMPLATE") or "/home/limour/pi-test/PPT模板.pptx"
        if not Path(template).exists():
            self.skipTest("real template not available")
        info = kit("inspect", template, "--slides", "1-3")
        self.assertEqual(len(info["slides"]), 3)
        deck = self.create("real.pptx", template=template)
        before = self.inspect(deck, "3")
        before_pics = [s for s in before["slides"][0]["shapes"] if s["type"] == "PICTURE"]
        kit("clone", str(deck), "3")
        after = self.inspect(deck, "4")
        after_pics = [s for s in after["slides"][0]["shapes"] if s["type"] == "PICTURE"]
        self.assertEqual(len(before_pics), len(after_pics))
        self.assertGreater(len(after_pics), 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
