"""The op layer: one function per mutation, shared by the CLI and ``apply``.

Every op takes the in-memory presentation plus a plain dict and returns a
JSON-able result. Ops never save; the caller commits once (atomic batch).
"""
from __future__ import annotations

from pathlib import Path

from pptx.util import Inches, Pt

from .shapes import resolve_shape, shape_brief
from .slides import clone_slide, get_layout, get_slide, remove_shape
from .tables import set_theme_table_style, style_table
from .text import (
    apply_para_spec,
    apply_run_spec,
    apply_text_frame_defaults,
    fill_paragraphs,
    write_paragraphs,
)
from .utils import (
    BULLET_NONE,
    MAX_CLONE_COUNT,
    VALIGN,
    box_in,
    fail,
    parse_enum,
    parse_pages,
    parse_rgb,
)

# Every op mutates the in-memory presentation and returns a JSON-able result.
# The CLI commands and the batch `apply` command both use this layer.
def op_set_text(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    shape, path = resolve_shape(slide, spec.get("shape"))
    if not shape.has_text_frame:
        if getattr(shape, "has_table", False) and shape.has_table:
            fail(f"{path} 是表格，请用 set_cell 修改单元格")
        fail(f"{path} 没有文本框，无法设置文本")
    tf = shape.text_frame
    if spec.get("paras") is not None:
        count = write_paragraphs(tf, spec["paras"])
    else:
        if spec.get("text") is None:
            fail("set_text 需要 text 或 paras")
        fill_paragraphs(tf, str(spec["text"]))
        count = len(tf.paragraphs)
    apply_name(shape, spec)
    # text="" clears the runs but keeps the shape; use delete_shape to remove it.
    return {
        "slide": int(spec["slide"]),
        "op": "set_text",
        "id": shape.shape_id,
        "shape": shape.name,
        "path": path,
        "text": tf.text,
        "paragraphs": count,
    }


def op_set_cell(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    shape, path = resolve_shape(slide, spec.get("shape"))
    if not (getattr(shape, "has_table", False) and shape.has_table):
        fail(f"{path} 不是表格")
    table = shape.table
    nrows, ncols = len(table.rows), len(table.columns)
    row, col = int(spec.get("row")), int(spec.get("col"))
    if not (0 <= row < nrows and 0 <= col < ncols):
        fail(f"单元格 ({row},{col}) 超出表格范围 {nrows}x{ncols}（0 起）")
    cell = table.cell(row, col)
    runs = spec.get("cell_runs", spec.get("cellRuns", spec.get("runs")))
    if runs is not None:
        from .tables import write_cell_runs

        write_cell_runs(cell, runs)
        return {
            "slide": int(spec["slide"]),
            "op": "set_cell",
            "id": shape.shape_id,
            "shape": shape.name,
            "path": path,
            "row": row,
            "col": col,
            "text": cell.text,
            "rich": True,
        }
    text = "" if spec.get("text") is None else str(spec["text"])
    old = cell.text_frame.paragraphs[0]
    if old.runs:
        old.runs[0].text = text
        for run in old.runs[1:]:
            run._r.getparent().remove(run._r)
        for para in list(cell.text_frame.paragraphs)[1:]:
            para._p.getparent().remove(para._p)
    else:
        cell.text = text
    return {
        "slide": int(spec["slide"]),
        "op": "set_cell",
        "id": shape.shape_id,
        "shape": shape.name,
        "path": path,
        "row": row,
        "col": col,
        "text": text,
    }


def apply_name(shape, spec: dict) -> None:
    """``name=`` gives shapes a stable handle for later ``ignore=`` / edits."""
    if spec.get("name"):
        shape.name = str(spec["name"])


PARA_KEYS = (
    "align",
    "space_after",
    "space_before",
    "line_spacing",
    "indent",
    "hanging",
    "bullet",
    "bullet_color",
    "level",
)

RUN_KEYS = ("size", "font", "color", "bold", "italic", "underline", "letter_spacing", "letterSpacing")


def write_text_spec(tf, spec: dict) -> int:
    """Fill a text frame from a flat (``text=``) or rich (``paras=``) spec."""
    para_defaults = {key: spec[key] for key in PARA_KEYS if spec.get(key) is not None}
    # A bullet with no explicit indent gets the usual 0.25" hanging layout
    # instead of drawing the glyph on top of the first character.
    if (
        para_defaults.get("bullet") not in BULLET_NONE
        and para_defaults.get("indent") is None
        and para_defaults.get("hanging") is None
    ):
        para_defaults["indent"] = 0.25
    # top-level size/font/color/bold apply to runs unless a run overrides them
    run_defaults = {key: spec[key] for key in RUN_KEYS if spec.get(key) is not None}
    if spec.get("paras") is not None:
        return write_paragraphs(tf, spec["paras"], para_defaults, run_defaults)
    text = "" if spec.get("text") is None else str(spec["text"])
    lines = text.split("\n")
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        apply_para_spec(para, para_defaults)
        run = para.add_run()
        run.text = line
        apply_run_spec(run, spec)
    return len(lines)


def op_add_text(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    x = float(spec["x"])
    y = float(spec["y"])
    w = float(spec.get("w", 6.0))
    fit = spec.get("fit")
    if spec.get("h") is not None:
        h = float(spec["h"])
    else:
        # No explicit height: grow with the text instead of a fixed 0.8" box
        # that clips three lines at 18pt.
        if fit is None:
            fit = "grow"
        size = float(spec.get("size") or 18)
        text = spec.get("text") or ""
        lines = max(1, len(str(text).split("\n")))
        if spec.get("paras"):
            lines = max(1, len(spec["paras"]))
        h = round(0.25 + lines * size * 1.35 / 72.0, 2)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    apply_text_frame_defaults(tf, {**spec, "fit": fit or "none"})
    count = write_text_spec(tf, spec)

    apply_name(box, spec)
    if spec.get("valign"):
        tf.vertical_anchor = parse_enum(spec["valign"], VALIGN, "垂直对齐")
    return {
        "slide": int(spec["slide"]),
        "op": "add_text",
        "added": "textbox",
        "id": box.shape_id,
        "name": box.name,
        "path": box.name,
        "box_in": box_in(box),
        "paragraphs": count,
        "fit": fit or "none",
    }


SHAPE_KINDS = {
    "rect": "RECTANGLE",
    "rectangle": "RECTANGLE",
    "roundrect": "ROUNDED_RECTANGLE",
    "round_rect": "ROUNDED_RECTANGLE",
    "rounded_rectangle": "ROUNDED_RECTANGLE",
    "chevron": "CHEVRON",
    "arrow": "RIGHT_ARROW",
    "rightarrow": "RIGHT_ARROW",
    "right_arrow": "RIGHT_ARROW",
    "leftarrow": "LEFT_ARROW",
    "left_arrow": "LEFT_ARROW",
    "uparrow": "UP_ARROW",
    "up_arrow": "UP_ARROW",
    "downarrow": "DOWN_ARROW",
    "down_arrow": "DOWN_ARROW",
    "oval": "OVAL",
    "circle": "OVAL",
    "triangle": "ISOSCELES_TRIANGLE",
    "parallelogram": "PARALLELOGRAM",
    "diamond": "DIAMOND",
    "pentagon": "PENTAGON",
    "hexagon": "HEXAGON",
    "bar": "RECTANGLE",
    "line": "LINE",
    "connector": "LINE",
}


def op_add_shape(prs, spec: dict) -> dict:
    """Add a drawn shape: card, accent bar, chevron, arrow, divider line.

    Covers what used to require hand-written python-pptx: rounded-corner radius,
    shadow suppression (``shadow=false`` default) and text inside the shape.
    """
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    slide = get_slide(prs, spec.get("slide"))
    kind = str(spec.get("kind") or "roundrect").strip().lower()
    if kind not in SHAPE_KINDS:
        fail(f"未知形状 kind={kind!r}。可选: {', '.join(sorted(SHAPE_KINDS))}")
    x = float(spec["x"])
    y = float(spec["y"])
    w = float(spec["w"])
    default_h = 0.06 if kind in ("line", "bar") else 0.6
    h = float(spec.get("h") if spec.get("h") is not None else default_h)
    is_line_kind = SHAPE_KINDS[kind] == "LINE"
    if w <= 0 or (h <= 0 and not is_line_kind):
        fail("add_shape 需要正的 w / h（分隔线 h 可为 0）")

    is_line = is_line_kind
    if is_line:
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y + h)
        )
    else:
        shape = slide.shapes.add_shape(
            getattr(MSO_SHAPE, SHAPE_KINDS[kind]), Inches(x), Inches(y), Inches(w), Inches(h)
        )

    if spec.get("shadow") is not None:
        shape.shadow.inherit = bool(spec["shadow"])
    else:
        shape.shadow.inherit = False

    fill = spec.get("fill")
    if is_line:
        if fill is None and spec.get("line") is None:
            fill = "#1F4E79"
        if fill:
            shape.line.color.rgb = parse_rgb(str(fill))
            shape.line.width = Pt(float(spec.get("line_width", spec.get("lineWidth", 1.0))))
        else:
            shape.line.fill.background()
    else:
        if fill in (None, "none", "false", False):
            if fill is None and spec.get("line") is None:
                fill = "#F4F7FB"
            if fill in (None, "none", "false", False):
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = parse_rgb(str(fill))
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_rgb(str(fill))
        line = spec.get("line")
        if line in (None,):
            line = "#CAD9E8"
        if line in ("none", "false", False):
            shape.line.fill.background()
        else:
            shape.line.color.rgb = parse_rgb(str(line))
            shape.line.width = Pt(float(spec.get("line_width", spec.get("lineWidth", 1.0))))

    radius = spec.get("radius", spec.get("corner_radius"))
    if radius is not None and not is_line:
        try:
            adjustments = shape.adjustments
            if len(adjustments):
                smaller = min(w, h)
                adjustments[0] = max(0.0, min(0.5, float(radius) / smaller)) if smaller else 0.0
        except Exception:  # presets without adjustments
            pass

    paragraphs = 0
    if spec.get("text") is not None or spec.get("paras") is not None:
        tf = shape.text_frame
        tf.word_wrap = True
        defaults = {"align": spec.get("align", "center"), "valign": spec.get("valign", "middle")}
        apply_text_frame_defaults(tf, {**defaults, **spec, "fit": spec.get("fit") or "none"})
        paragraphs = write_text_spec(tf, spec)

    apply_name(shape, spec)
    return {
        "slide": int(spec["slide"]),
        "op": "add_shape",
        "added": "shape",
        "kind": kind,
        "id": shape.shape_id,
        "name": shape.name,
        "path": shape.name,
        "box_in": box_in(shape),
        "paragraphs": paragraphs,
    }


def op_add_image(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    image_path = Path(str(spec.get("path"))).expanduser()
    if not image_path.is_absolute():
        image_path = (Path.cwd() / image_path).resolve()
    if not image_path.exists():
        fail(f"图片不存在: {image_path}")
    kwargs = {}
    if spec.get("x") is not None and spec.get("y") is not None:
        kwargs["left"] = Inches(float(spec["x"]))
        kwargs["top"] = Inches(float(spec["y"]))
    if spec.get("w") is not None:
        kwargs["width"] = Inches(float(spec["w"]))
    if spec.get("h") is not None:
        kwargs["height"] = Inches(float(spec["h"]))
    picture = slide.shapes.add_picture(str(image_path), **kwargs)
    apply_name(picture, spec)
    return {
        "slide": int(spec["slide"]),
        "op": "add_image",
        "added": "picture",
        "id": picture.shape_id,
        "name": picture.name,
        "path": picture.name,
        "box_in": box_in(picture),
    }


def op_add_table(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    rows, cols = int(spec.get("rows")), int(spec.get("cols"))
    if rows < 1 or cols < 1:
        fail("表格至少 1 行 1 列")
    data = spec.get("data") or []
    if not isinstance(data, list) or any(not isinstance(row, list) for row in data):
        fail("data 必须是二维数组")
    if len(data) > rows or any(len(row) > cols for row in data):
        fail(f"data 尺寸 {len(data)}x{max((len(r) for r in data), default=0)} 超过 {rows}x{cols}")

    had_table = any(
        getattr(shape, "has_table", False) and shape.has_table
        for slide_obj in prs.slides
        for shape in slide_obj.shapes
    )
    frame = slide.shapes.add_table(
        rows,
        cols,
        Inches(float(spec["x"])),
        Inches(float(spec["y"])),
        Inches(float(spec["w"])),
        Inches(float(spec["h"])),
    )
    table = frame.table
    theme_style = bool(spec.get("theme_style"))
    set_theme_table_style(frame, theme_style)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            value = ""
            if r < len(data) and c < len(data[r]):
                value = "" if data[r][c] is None else str(data[r][c])
            cell.text = value
    border = spec.get("border", "#D9D9D9")
    if border is None:
        border = "#D9D9D9"
    style_spec = {
        "header": bool(spec.get("header", True)),
        "header_fill": spec.get("header_fill"),
        "header_color": spec.get("header_color"),
        "body_color": spec.get("body_color"),
        "zebra": spec.get("zebra"),
        "border": border,
        "body_size": spec.get("size"),
        "header_size": spec.get("size"),
        "col_widths": spec.get("col_widths"),
        "row_height": spec.get("row_height"),
        "aligns": spec.get("aligns"),
        "margin": spec.get("margin"),
        "valign": spec.get("valign"),
        "header_bold": spec.get("header", True),
        "header_height": spec.get("headerHeight", spec.get("header_height")),
        "first_col_bold": spec.get("first_col_bold", spec.get("firstColBold")),
        "first_col_color": spec.get("first_col_color", spec.get("firstColColor")),
        "header_border": spec.get("header_border", spec.get("headerBorder")),
        "cell_runs": spec.get("cell_runs", spec.get("cellRuns")),
    }
    style_table(frame, style_spec)
    apply_name(frame, spec)
    result = {
        "slide": int(spec["slide"]),
        "op": "add_table",
        "added": "table",
        "id": frame.shape_id,
        "name": frame.name,
        "path": frame.name,
        "box_in": box_in(frame),
        "rows": rows,
        "cols": cols,
        "theme_style": theme_style,
    }
    if not had_table:
        result["hint"] = (
            "模板里没有任何表格样例：已关闭主题表样式并使用中性边框。"
            "可用 ppt_table_style 调整，或克隆带表格的页。"
        )
    return result


def _move_in_z_order(slide, shape, mode: str) -> None:
    """Reorder a shape inside the slide's spTree (front/back/forward/backward)."""
    spTree = slide.shapes._spTree
    element = shape._element
    parent = element.getparent()
    if parent is None or parent is not spTree:
        fail("只能调整顶层形状的叠放次序（组合内请先选中组合）")
    children = [child for child in spTree if child.tag.endswith("}sp") or child.tag.endswith("}pic")
                or child.tag.endswith("}graphicFrame") or child.tag.endswith("}grpSp")
                or child.tag.endswith("}cxnSp")]
    if element not in children:
        fail("这个形状不在可排序的绘图元素里")
    index = children.index(element)
    if mode == "front":
        spTree.remove(element)
        spTree.append(element)
        return
    if mode == "back":
        spTree.remove(element)
        first = children[0]
        first.addprevious(element)
        return
    if mode == "forward":
        if index + 1 < len(children):
            children[index + 1].addnext(element)
        return
    if mode == "backward":
        if index > 0:
            children[index - 1].addprevious(element)
        return
    fail(f"未知的叠放次序 z={mode!r}（front|back|forward|backward）")


def op_set_shape(prs, spec: dict) -> dict:
    """Move/resize a shape, restyle its fill/border, rename it or change z-order.

    Templates place placeholders where the designer wanted them; a page that
    reuses a layout still needs to move its title. Without this the only way out
    was deleting the placeholder and drawing a new box (losing the semantics).
    """
    slide = get_slide(prs, spec.get("slide"))
    shape, path = resolve_shape(slide, spec.get("shape"))
    before = box_in(shape)
    moved = {}
    if spec.get("x") is not None:
        shape.left = Inches(float(spec["x"]))
    if spec.get("y") is not None:
        shape.top = Inches(float(spec["y"]))
    if spec.get("w") is not None:
        shape.width = Inches(float(spec["w"]))
    if spec.get("h") is not None:
        shape.height = Inches(float(spec["h"]))
    if any(spec.get(k) is not None for k in ("x", "y", "w", "h")):
        moved["box_in"] = box_in(shape)

    fill = spec.get("fill")
    if fill is not None:
        if fill in ("none", False, "false"):
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_rgb(str(fill))
        moved["fill"] = fill
    line = spec.get("line")
    if line is not None:
        if line in ("none", False, "false"):
            shape.line.fill.background()
        else:
            shape.line.color.rgb = parse_rgb(str(line))
            if spec.get("line_width") is not None:
                shape.line.width = Pt(float(spec["line_width"]))
        moved["line"] = line
    if spec.get("rotation") is not None:
        shape.rotation = float(spec["rotation"])
        moved["rotation"] = shape.rotation
    apply_name(shape, spec)
    if spec.get("z"):
        _move_in_z_order(slide, shape, str(spec["z"]).lower())
        moved["z"] = str(spec["z"]).lower()

    return {
        "slide": int(spec["slide"]),
        "op": "set_shape",
        "id": shape.shape_id,
        "name": shape.name,
        "path": path,
        "box_before": before,
        **moved,
    }


def op_delete_shape(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    shape, path = resolve_shape(slide, spec.get("shape"))
    info = {
        "slide": int(spec["slide"]),
        "op": "delete_shape",
        "id": shape.shape_id,
        "shape": shape.name,
        "path": path,
        "deleted": True,
    }
    remove_shape(shape)
    return info


def op_set_notes(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    text = "" if spec.get("text") is None else str(spec["text"])
    tf = slide.notes_slide.notes_text_frame
    tf.text = text
    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            title = shape.text_frame.text.strip().split("\n")[0]
            if "标题" in shape.name.lower() or "title" in shape.name.lower():
                break
    return {
        "slide": int(spec["slide"]),
        "op": "set_notes",
        "chars": len(text),
        "paragraphs": len(tf.paragraphs),
        "first_20": text.strip()[:20],
        "slide_title": title[:20],
    }


def op_table_style(prs, spec: dict) -> dict:
    slide = get_slide(prs, spec.get("slide"))
    shape, path = resolve_shape(slide, spec.get("shape"))
    if not (getattr(shape, "has_table", False) and shape.has_table):
        fail(f"{path} 不是表格")
    info = style_table(shape, spec)
    return {
        "slide": int(spec["slide"]),
        "op": "table_style",
        "id": shape.shape_id,
        "shape": shape.name,
        "path": path,
        **info,
    }


def op_add_slide(prs, spec: dict) -> dict:
    layout = get_layout(prs, spec.get("layout"))
    slide = prs.slides.add_slide(layout)
    return {
        "slide": len(prs.slides),
        "op": "add_slide",
        "layout": layout.name,
        "placeholders": [ph.name for ph in slide.placeholders],
        "shapes": [shape_brief(shape, i) for i, shape in enumerate(slide.shapes)],
    }


def op_clone(prs, spec: dict) -> dict:
    total = len(prs.slides)
    src = int(spec.get("source", spec.get("src", 0)))
    if not 1 <= src <= total:
        fail(f"源页号 {src} 超出 1..{total}")
    count = int(spec.get("count") or 1)
    if not 1 <= count <= MAX_CLONE_COUNT:
        fail(f"克隆数量应在 1..{MAX_CLONE_COUNT}，收到 {count}")
    at = spec.get("at")
    if at is not None and not 1 <= int(at) <= total + 1:
        fail(f"插入位置 {at} 超出 1..{total + 1}")
    clone_slide(prs, src - 1, int(at) if at is not None else None, count)
    return {
        "slide": int(at) if at is not None else len(prs.slides),
        "op": "clone",
        "created_from": src,
        "count": count,
        "at": at,
        "slide_count": len(prs.slides),
    }


def op_delete_slides(prs, spec: dict) -> dict:
    total = len(prs.slides)
    pages = parse_pages(str(spec.get("slides")), total)
    if len(pages) >= total:
        fail("不能删除全部页（至少保留 1 页）")
    sld_id_lst = prs.slides._sldIdLst
    ids = list(sld_id_lst)
    for page in sorted(pages, reverse=True):
        sld_id_lst.remove(ids[page - 1])
    notes = []
    for idx, slide in enumerate(prs.slides, start=1):
        if not slide.has_notes_slide:
            continue
        text = slide.notes_slide.notes_text_frame.text.strip()
        if text:
            notes.append({"slide": idx, "notes_len": len(text), "first_20": text[:20]})
    result = {
        "slide": pages[0],
        "op": "delete_slides",
        "deleted": pages,
        "slide_count": len(prs.slides),
    }
    if notes:
        result["notes_after"] = notes
        result["notes_hint"] = "删除页后请核对备注是否仍对应当前页（notes_check 可复核）。"
    return result


OP_HANDLERS = {
    "set_text": op_set_text,
    "set_cell": op_set_cell,
    "add_text": op_add_text,
    "add_shape": op_add_shape,
    "add_image": op_add_image,
    "add_table": op_add_table,
    "set_shape": op_set_shape,
    "delete_shape": op_delete_shape,
    "set_notes": op_set_notes,
    "table_style": op_table_style,
    "add_slide": op_add_slide,
    "clone": op_clone,
    "delete_slides": op_delete_slides,
}


def apply_one(prs, op) -> dict:
    if not isinstance(op, dict):
        fail("每个 op 必须是 JSON 对象")
    name = op.get("op") or op.get("type")
    handler = OP_HANDLERS.get(str(name))
    if handler is None:
        fail(f"未知 op {name!r}。支持: {', '.join(sorted(OP_HANDLERS))}")
    result = handler(prs, op)
    result.setdefault("op", name)
    return result
