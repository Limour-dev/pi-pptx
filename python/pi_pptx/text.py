"""Run/paragraph writing: fonts, bullets, indents, rich text (``paras=``).

Also home to :func:`text_fit`, the shared "does this text still fit its box"
estimator used by ``ppt_inspect`` (before drawing) and ``ppt_shot`` (density).
"""
from __future__ import annotations

import copy

from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .utils import ALIGN, BULLET_NONE, EMU, FIT, PT_PER_IN, VALIGN, fail, parse_enum, parse_rgb


def _set_run_font(run, font: str) -> None:
    """Set latin *and* East-Asian/complex-script typefaces.

    python-pptx's ``run.font.name`` only writes ``a:latin``; CJK glyphs follow
    ``a:ea`` (or the theme), so ``font='微软雅黑'`` used to be a silent no-op for
    Chinese text.
    """
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    prev = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            if prev is not None:
                prev.addnext(el)
            else:
                rPr.append(el)
        el.set("typeface", font)
        prev = el


def _first_run_rpr(text_frame):
    for para in text_frame.paragraphs:
        for run in para.runs:
            rPr = run._r.rPr
            if rPr is not None:
                return copy.deepcopy(rPr)
            return None
    return None


def apply_run_spec(run, spec: dict) -> None:
    if spec.get("size") is not None:
        run.font.size = Pt(float(spec["size"]))
    if spec.get("font"):
        _set_run_font(run, str(spec["font"]))
    if spec.get("color"):
        run.font.color.rgb = parse_rgb(str(spec["color"]))
    if spec.get("bold") is not None:
        run.font.bold = bool(spec["bold"])
    if spec.get("italic") is not None:
        run.font.italic = bool(spec["italic"])
    if spec.get("underline") is not None:
        run.font.underline = bool(spec["underline"])
    spacing = spec.get("letter_spacing", spec.get("letterSpacing", spec.get("spacing")))
    if spacing is not None:
        set_letter_spacing(run, spacing)


def set_letter_spacing(run, value) -> None:
    """Write ``a:rPr/@spc`` (1/100 pt). python-pptx has no API for this.

    ``value`` is in points (``0.5`` = half a point of tracking); a trailing
    ``em`` (``"0.02em"``) is converted against the run's font size.
    """
    rPr = run._r.get_or_add_rPr()
    if isinstance(value, str) and value.strip().lower().endswith("em"):
        em = float(value.strip()[:-2])
        size = rPr.get("sz")
        size_pt = float(size) / 100.0 if size else 18.0
        hundredths = int(round(em * size_pt * 100))
    else:
        hundredths = int(round(float(value) * 100))
    rPr.set("spc", str(hundredths))


def _bullet_anchor(pPr):
    for tag in ("a:tabLst", "a:defRPr", "a:extLst"):
        el = pPr.find(qn(tag))
        if el is not None:
            return el
    return None


def set_paragraph_bullet(
    para,
    char=None,
    color: str | None = None,
    indent: float | None = None,
    hanging: float | None = None,
) -> None:
    """Write a bullet (or explicit buNone) plus hanging indent into a:pPr.

    ``char="__keep__"`` only updates the indents and leaves the existing bullet
    definition alone (used when a paragraph is given an indent but no bullet).
    """
    pPr = para._p.get_or_add_pPr()
    keep = char == "__keep__"
    if not keep:
        for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)

    if indent is not None:
        mar = abs(float(indent))
        pPr.set("marL", str(int(Inches(mar))))
        if hanging is None:
            hanging = mar if char not in BULLET_NONE and not keep else 0.0
        pPr.set("indent", str(-int(Inches(max(0.0, float(hanging))))))
    elif hanging is not None:
        pPr.set("indent", str(-int(Inches(max(0.0, float(hanging))))))
    if keep:
        return

    anchor = _bullet_anchor(pPr)
    elements = []
    if char in BULLET_NONE:
        elements.append(pPr.makeelement(qn("a:buNone"), {}))
    else:
        glyph = "•" if char is True else str(char)
        if color:
            buClr = pPr.makeelement(qn("a:buClr"), {})
            buClr.append(buClr.makeelement(qn("a:srgbClr"), {"val": str(parse_rgb(str(color)))}))
            elements.append(buClr)
        buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
        elements.append(buFont)
        elements.append(pPr.makeelement(qn("a:buChar"), {"char": glyph}))
    for el in elements:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            pPr.append(el)


def apply_para_spec(para, spec: dict) -> None:
    align = parse_enum(spec.get("align"), ALIGN, "对齐方式")
    if align is not None:
        para.alignment = align
    if spec.get("space_after") is not None:
        para.space_after = Pt(float(spec["space_after"]))
    if spec.get("space_before") is not None:
        para.space_before = Pt(float(spec["space_before"]))
    if spec.get("line_spacing") is not None:
        para.line_spacing = float(spec["line_spacing"])
    if spec.get("level") is not None:
        para.level = int(spec["level"])
    bullet_given = "bullet" in spec
    if bullet_given or spec.get("indent") is not None or spec.get("hanging") is not None:
        set_paragraph_bullet(
            para,
            char=spec.get("bullet") if bullet_given else "__keep__",
            color=spec.get("bullet_color"),
            indent=spec.get("indent"),
            hanging=spec.get("hanging"),
        )


def _write_runs(para, runs: list[dict], base_rpr) -> None:
    for run in list(para.runs):
        run._r.getparent().remove(run._r)
    for spec in runs:
        if not isinstance(spec, dict):
            spec = {"t": spec}
        run = para.add_run()
        run.text = "" if spec.get("t", spec.get("text", "")) is None else str(spec.get("t", spec.get("text", "")))
        if base_rpr is not None:
            rPr = run._r.get_or_add_rPr()
            run._r.remove(rPr)
            run._r.insert(0, copy.deepcopy(base_rpr))
        apply_run_spec(run, spec)


def write_paragraphs(
    text_frame,
    paras: list[dict],
    defaults: dict | None = None,
    run_defaults: dict | None = None,
) -> int:
    """Replace all paragraphs with run-level rich text (the `paras=` form)."""
    if not isinstance(paras, list):
        fail("paras 必须是数组，每项是一段")
    defaults = defaults or {}
    run_defaults = run_defaults or {}
    base_rpr = _first_run_rpr(text_frame)
    base_p = copy.deepcopy(text_frame.paragraphs[0]._p) if text_frame.paragraphs else None
    for para in list(text_frame.paragraphs):
        para._p.getparent().remove(para._p)
    if not paras:
        paras = [{"runs": [{"t": ""}]}]
    for spec in paras:
        if not isinstance(spec, dict):
            fail("paras 的每一项必须是对象")
        el = copy.deepcopy(base_p) if base_p is not None else text_frame._txBody.makeelement(qn("a:p"), {})
        text_frame._txBody.append(el)
        para = text_frame.paragraphs[-1]
        for run in list(para.runs):
            run._r.getparent().remove(run._r)
        merged = {**defaults, **spec}
        apply_para_spec(para, merged)
        runs = spec.get("runs")
        if runs is None:
            runs = [{"t": spec.get("t", spec.get("text", ""))}]
        if not isinstance(runs, list):
            fail("paras[].runs 必须是数组")
        if run_defaults:
            runs = [{**run_defaults, **run} if isinstance(run, dict) else run for run in runs]
        _write_runs(para, runs, base_rpr)
    return len(text_frame.paragraphs)


def _set_para_text(para, text: str) -> None:
    """Replace a paragraph's text while keeping the first run's formatting."""
    runs = list(para.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        para.add_run().text = text


def fill_paragraphs(text_frame, text: str) -> None:
    """Write text (\\n separated) inheriting the first paragraph's style."""
    lines = text.split("\n")
    first = text_frame.paragraphs[0]
    base = copy.deepcopy(first._p)
    for para in list(text_frame.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    _set_para_text(first, lines[0])
    prev = first._p
    for _ in lines[1:]:
        element = copy.deepcopy(base)
        prev.addnext(element)
        prev = element
    for para, line in zip(text_frame.paragraphs, lines):
        _set_para_text(para, line)


def apply_text_frame_defaults(text_frame, spec: dict) -> None:
    fit = spec.get("fit")
    if fit is not None:
        text_frame.auto_size = parse_enum(fit, FIT, "fit")
    if spec.get("valign"):
        text_frame.vertical_anchor = parse_enum(spec["valign"], VALIGN, "垂直对齐")
    if spec.get("wrap") is False:
        text_frame.word_wrap = False


# ------------------------------------------------------------------ text fit
# Glyph width in em: CJK/full-width are 1, Latin averages ~0.52. A real font
# metric would need the font file; this is the same ballpark the layout heuristics
# need, and it errs on the *wide* side (5-10% margin for font substitution).
_FULLWIDTH_RANGES = (
    (0x1100, 0x115F),  # Hangul Jamo
    (0x2E80, 0xA4CF),  # CJK radicals .. Yi
    (0xAC00, 0xD7A3),  # Hangul syllables
    (0xF900, 0xFAFF),  # CJK compatibility ideographs
    (0xFE30, 0xFE6F),  # CJK compatibility forms
    (0xFF00, 0xFF60),  # full-width forms
    (0xFFE0, 0xFFE6),
    (0x1F300, 0x1FAFF),  # emoji
)


def _char_em(char: str) -> float:
    code = ord(char)
    for low, high in _FULLWIDTH_RANGES:
        if low <= code <= high:
            return 1.0
    if char.isspace():
        return 0.28
    if not char.isprintable():
        return 0.0
    return 0.52


def _atoms(text: str) -> list[str]:
    """Split into wrap atoms: one per CJK glyph, one per Latin word.

    PowerPoint breaks CJK anywhere but Latin only at spaces, so a single wide
    token ("01") overflows sideways instead of wrapping — treating it as one
    atom keeps the line estimate honest for big numbers in tight boxes.
    """
    atoms: list[str] = []
    buffer = ""
    for char in text:
        if _char_em(char) >= 1.0 or char.isspace():
            if buffer:
                atoms.append(buffer)
                buffer = ""
            atoms.append(" " if char.isspace() else char)
        else:
            buffer += char
    if buffer:
        atoms.append(buffer)
    return atoms


def _wrap_lines(text: str, size: float, usable_w_pt: float) -> int:
    if not text.strip():
        return 1
    lines = 1
    current = 0.0
    for atom in _atoms(text):
        if atom == " ":
            current += 0.28 * size
            continue
        width = sum(_char_em(char) for char in atom) * size
        if current > 0 and current + width > usable_w_pt:
            lines += 1
            current = width
        else:
            current += width
    return lines


def _frame_margins(text_frame) -> tuple[float, float, float, float]:
    """left, right, top, bottom insets in EMU (DrawingML defaults: 0.1in / 0.05in)."""
    def value(name, default):
        inset = getattr(text_frame, name, None)
        return default if inset is None else int(inset)

    return (
        value("margin_left", 91440),
        value("margin_right", 91440),
        value("margin_top", 45720),
        value("margin_bottom", 45720),
    )


def text_fit(shape) -> dict | None:
    """Estimate whether a shape's text still fits the shape's box.

    Answers the question the agent could not answer before rendering: "how many
    lines does this box need, and how many does it hold?" ``needed_in`` is
    ``lines x size x 1.35`` (the rule of thumb a hand-built deck uses) plus the
    text-frame insets, so it is comparable with ``box_h_in``.

    ``fits`` compares *line counts*, not inches: a 13pt single line in a 0.28in
    label box spills a few points of leading by geometry yet renders exactly as
    intended, which is why the inch comparison alone cried wolf on every label.
    ``slack_in`` is still reported so a caller can see the raw margin.

    Returns ``None`` for shapes without text, otherwise ``size_pt``,
    ``est_lines``, ``capacity_lines``, ``needed_in``, ``box_h_in``, ``slack_in``
    and ``fits``.
    """
    if not getattr(shape, "has_text_frame", False):
        return None
    tf = shape.text_frame
    if shape.width is None or shape.height is None:
        return None
    if not tf.text.strip():
        return None
    margin_l, margin_r, margin_t, margin_b = _frame_margins(tf)
    usable_w_pt = max(1.0, (shape.width - margin_l - margin_r) / EMU * PT_PER_IN)
    box_h_in = round(shape.height / EMU, 3)

    lines = 0
    needed_pt = 0.0
    max_size = 0.0
    line_h_in = 0.0
    for para in tf.paragraphs:
        size = next((run.font.size.pt for run in para.runs if run.font.size is not None), None)
        size = float(size) if size is not None else 18.0
        max_size = max(max_size, size)
        spacing = 1.0
        if para.line_spacing is not None and isinstance(para.line_spacing, float):
            spacing = max(0.5, float(para.line_spacing))
        text = "".join(run.text for run in para.runs) or para.text
        para_lines = _wrap_lines(text, size, usable_w_pt)
        lines += para_lines
        space_after = para.space_after.pt if para.space_after is not None else 0.0
        needed_pt += para_lines * size * 1.35 * spacing + float(space_after)
        line_h_in = max(line_h_in, size * 1.35 * spacing / PT_PER_IN)
    needed_in = round(needed_pt / PT_PER_IN + (margin_t + margin_b) / EMU, 3)
    inset_in = (margin_t + margin_b) / EMU
    capacity = max(1, int((box_h_in - inset_in) // line_h_in)) if line_h_in else 1
    return {
        "estimated": True,
        "size_pt": round(max_size, 1),
        "est_lines": lines,
        "capacity_lines": capacity,
        "needed_in": needed_in,
        "box_h_in": box_h_in,
        "slack_in": round(box_h_in - needed_in, 3),
        "fits": lines <= capacity,
    }
