"""Slide helpers: cloning (with relationship remapping), lookup, deletion."""
from __future__ import annotations

import copy

from pptx.oxml.ns import qn

from .utils import R_NS, fail

def remap_rels(dest_slide, src_slide) -> None:
    """Re-point every r:id / r:embed in copied XML at the destination rels.

    Without this step pictures, charts and links silently disappear.
    """
    dpart, spart = dest_slide.part, src_slide.part
    for element in dest_slide.element.iter():
        for attr, value in list(element.attrib.items()):
            if not attr.startswith("{%s}" % R_NS):
                continue
            rel = spart.rels.get(value)
            if rel is None:
                continue
            if rel.is_external:
                element.attrib[attr] = dpart.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            else:
                element.attrib[attr] = dpart.relate_to(rel.target_part, rel.reltype)


def clone_slide(prs, src_index: int, at: int | None = None, count: int = 1) -> list:
    src = prs.slides[src_index]
    created = []
    for k in range(count):
        new = prs.slides.add_slide(src.slide_layout)
        for shape in list(new.shapes):  # drop placeholders inherited from the layout
            shape._element.getparent().remove(shape._element)
        for shape in src.shapes:
            new.shapes._spTree.append(copy.deepcopy(shape._element))

        # copy explicit slide background (if any) so themed pages stay themed
        src_bg = src.element.cSld.find(qn("p:bg"))
        dest_bg = new.element.cSld.find(qn("p:bg"))
        if src_bg is not None:
            if dest_bg is not None:
                new.element.cSld.replace(dest_bg, copy.deepcopy(src_bg))
            else:
                new.element.cSld.insert(0, copy.deepcopy(src_bg))
        elif dest_bg is not None:
            new.element.cSld.remove(dest_bg)

        remap_rels(new, src)

        if src.has_notes_slide:
            notes = src.notes_slide.notes_text_frame.text
            if notes.strip():
                new.notes_slide.notes_text_frame.text = notes

        created.append(new)
        if at is not None:
            sld_id_lst = prs.slides._sldIdLst
            ids = list(sld_id_lst)
            moved = ids[-1]
            sld_id_lst.remove(moved)
            pos = min(max(at - 1 + k, 0), len(ids) - 1)
            sld_id_lst.insert(pos, moved)
    return created


def get_slide(prs, number) -> "object":
    total = len(prs.slides)
    try:
        number = int(number)
    except (TypeError, ValueError):
        fail(f"页号必须是整数，收到: {number!r}")
    if not 1 <= number <= total:
        fail(f"页号 {number} 超出 1..{total}")
    return prs.slides[number - 1]


def get_layout(prs, key):
    key = str(key).strip()
    if key.lstrip("-").isdigit():
        idx = int(key)
        if not 0 <= idx < len(prs.slide_layouts):
            fail(f"版式序号 {idx} 超出 0..{len(prs.slide_layouts) - 1}")
        return prs.slide_layouts[idx]
    for cand in prs.slide_layouts:
        if cand.name == key:
            return cand
    names = [f"{i}:{l.name}" for i, l in enumerate(prs.slide_layouts)]
    fail(f"找不到版式 {key!r}。可用版式: {names}")


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)
