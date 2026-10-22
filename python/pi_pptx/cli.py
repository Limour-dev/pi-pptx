"""Command-line interface: one subcommand per backend operation.

``python3 python/pptkit.py <cmd> ...`` still works unchanged; this module is
also reachable as ``python3 -m pi_pptx.cli``.
"""
from __future__ import annotations

import argparse
import os
import shutil
import sys
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from .doctor import cmd_doctor
from .ops import (
    apply_one,
    op_add_image,
    op_add_shape,
    op_add_slide,
    op_add_table,
    op_add_text,
    op_clone,
    op_delete_shape,
    op_delete_slides,
    op_set_cell,
    op_set_notes,
    op_set_shape,
    op_set_text,
    op_table_style,
)
from .render import deck_pdf, preview_overflow, render_overflow, shot_deck
from .shapes import inspect_deck, notes_check
from .utils import (
    KitError,
    deck_meta,
    emit,
    fail,
    load,
    parse_json_arg,
    save_and_reload,
    save_deck,
)

def cmd_inspect(args):
    prs, path = load(args.pptx)
    emit(inspect_deck(prs, path, mode=args.mode, slides=args.slides, depth=args.depth))


def cmd_shot(args):
    prs, path = load(args.pptx)
    payload = shot_deck(
        prs,
        path,
        slides=args.slides,
        dpi=args.dpi,
        out=args.out,
        max_pages=args.max_pages,
        montage=args.montage,
        region=args.region,
        overflow=args.overflow,
        ignore=args.ignore,
        only_changed=args.only_changed,
        changed_slides=args.changed_slides,
        density=args.density,
    )
    emit(payload)


def cmd_notes_check(args):
    prs, _path = load(args.pptx)
    emit(notes_check(prs, args.slides))


def cmd_add_shape(args):
    prs, path = load(args.pptx)
    spec = {
        "slide": args.slide,
        "kind": args.kind,
        "x": args.x,
        "y": args.y,
        "w": args.w,
        "h": args.h,
        "radius": args.radius,
        "fill": args.fill,
        "line": args.line,
        "line_width": args.line_width,
        "shadow": args.shadow,
        "text": args.text,
        "size": args.size,
        "color": args.color,
        "bold": args.bold,
        "align": args.align,
        "valign": args.valign,
        "font": args.font,
        "name": args.name,
        "letter_spacing": args.letter_spacing,
    }
    if args.paras:
        spec["paras"] = parse_json_arg(args.paras, "--paras")
        spec.pop("text", None)
    result = op_add_shape(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_create(args):
    target = Path(args.out).expanduser()
    if not target.is_absolute():
        target = (Path.cwd() / target).resolve()
    if target.exists() and not args.overwrite:
        fail(f"目标文件已存在: {target}（需要覆盖请传 overwrite=true）")
    target.parent.mkdir(parents=True, exist_ok=True)

    if args.template:
        src = Path(args.template).expanduser()
        if not src.exists():
            fail(f"模板不存在: {src}")
        if src.resolve() == target.resolve():
            fail("template 与目标文件相同：请指定新的输出路径，不要覆盖模板")
        Presentation(str(src))  # verify it parses
        if args.width is None and args.height is None:
            tmp = target.parent / f".{target.name}.{uuid.uuid4().hex}.tmp"
            try:
                shutil.copyfile(src, tmp)
                Presentation(str(tmp))  # integrity
                os.replace(tmp, target)
            finally:
                if tmp.exists():
                    tmp.unlink()
        else:
            prs = Presentation(str(src))
            if args.width:
                prs.slide_width = Inches(args.width)
            if args.height:
                prs.slide_height = Inches(args.height)
            save_deck(prs, target, str(target))
        Presentation(str(target))  # integrity
    else:
        prs = Presentation()
        prs.slide_width = Inches(args.width or 13.333)
        prs.slide_height = Inches(args.height or 7.5)
        if not args.empty:
            layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(layout)
            if len(slide.placeholders) > 0:
                body = slide.placeholders[0]
                if body.has_text_frame and not body.text_frame.text.strip():
                    body.text_frame.text = "标题"
            if len(slide.placeholders) > 1:
                sub = slide.placeholders[1]
                if sub.has_text_frame and not sub.text_frame.text.strip():
                    sub.text_frame.text = "副标题"
        prs.save(str(target))
        Presentation(str(target))

    prs = Presentation(str(target))
    out = deck_meta(prs, target)
    out["template"] = str(Path(args.template).expanduser().resolve()) if args.template else None
    emit(out)


def cmd_add_slide(args):
    prs, path = load(args.pptx)
    result = op_add_slide(prs, {"layout": args.layout})
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    out["added_slide"] = result["slide"]
    emit(out)


def cmd_delete(args):
    prs, path = load(args.pptx)
    result = op_delete_slides(prs, {"slides": args.slides})
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_clone(args):
    prs, path = load(args.pptx)
    result = op_clone(
        prs,
        {"source": args.src, "at": args.at, "count": args.count or 1},
    )
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_set_text(args):
    prs, path = load(args.pptx)
    spec = {"slide": args.slide, "shape": args.shape}
    if getattr(args, "name", None):
        spec["name"] = args.name
    if args.paras:
        spec["paras"] = parse_json_arg(args.paras, "--paras")
    else:
        spec["text"] = args.text
    result = op_set_text(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_set_cell(args):
    prs, path = load(args.pptx)
    spec = {"slide": args.slide, "shape": args.shape, "row": args.row, "col": args.col}
    if getattr(args, "cell_runs", None):
        spec["cell_runs"] = parse_json_arg(args.cell_runs, "--cell-runs")
    else:
        spec["text"] = args.text
    result = op_set_cell(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_add_text(args):
    prs, path = load(args.pptx)
    spec = {
        "slide": args.slide,
        "text": args.text,
        "x": args.x,
        "y": args.y,
        "w": args.w,
        "h": args.h,
        "size": args.size,
        "font": args.font,
        "color": args.color,
        "bold": args.bold,
        "align": args.align,
        "valign": args.valign,
        "line_spacing": args.line_spacing,
        "space_after": args.space_after,
        "space_before": args.space_before,
        "indent": args.indent,
        "hanging": args.hanging,
        "bullet": args.bullet,
        "bullet_color": args.bullet_color,
        "letter_spacing": getattr(args, "letter_spacing", None),
        "name": getattr(args, "name", None),
        "fit": args.fit,
    }
    if args.paras:
        spec["paras"] = parse_json_arg(args.paras, "--paras")
        spec.pop("text", None)
    result = op_add_text(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_add_image(args):
    prs, path = load(args.pptx)
    result = op_add_image(
        prs,
        {
            "slide": args.slide,
            "path": args.path,
            "x": args.x,
            "y": args.y,
            "w": args.w,
            "h": args.h,
        },
    )
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_add_table(args):
    prs, path = load(args.pptx)
    spec = {
        "slide": args.slide,
        "rows": args.rows,
        "cols": args.cols,
        "x": args.x,
        "y": args.y,
        "w": args.w,
        "h": args.h,
        "size": args.size,
        "header": args.header,
        "theme_style": args.theme_style,
        "header_fill": args.header_fill,
        "header_color": args.header_color,
        "zebra": args.zebra,
        "border": args.border,
        "col_widths": args.col_widths,
        "row_height": args.row_height,
        "aligns": args.aligns,
    }
    if getattr(args, "header_height", None) is not None:
        spec["header_height"] = args.header_height
    if getattr(args, "first_col_bold", None) is not None:
        spec["first_col_bold"] = args.first_col_bold
    if getattr(args, "first_col_color", None):
        spec["first_col_color"] = args.first_col_color
    if getattr(args, "name", None):
        spec["name"] = args.name
    if getattr(args, "cell_runs", None):
        spec["cell_runs"] = parse_json_arg(args.cell_runs, "--cell-runs")
    if args.data:
        data = parse_json_arg(args.data, "--data")
        spec["data"] = data
    result = op_add_table(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_set_shape(args):
    prs, path = load(args.pptx)
    spec = {
        "slide": args.slide,
        "shape": args.shape,
        "x": args.x,
        "y": args.y,
        "w": args.w,
        "h": args.h,
        "fill": args.fill,
        "line": args.line,
        "line_width": args.line_width,
        "rotation": args.rotation,
        "z": args.z,
        "name": args.name,
    }
    result = op_set_shape(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_delete_shape(args):
    prs, path = load(args.pptx)
    result = op_delete_shape(prs, {"slide": args.slide, "shape": args.shape})
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_set_notes(args):
    prs, path = load(args.pptx)
    result = op_set_notes(prs, {"slide": args.slide, "text": args.text})
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_table_style(args):
    prs, path = load(args.pptx)
    spec = {
        "slide": args.slide,
        "shape": args.shape,
        "header": args.header,
        "header_fill": args.header_fill,
        "header_color": args.header_color,
        "body_color": args.body_color,
        "zebra": args.zebra,
        "border": args.border,
        "border_width": args.border_width,
        "body_size": args.body_size,
        "header_size": args.header_size,
        "header_bold": args.header_bold,
        "body_bold": args.body_bold,
        "col_widths": args.col_widths,
        "row_height": args.row_height,
        "aligns": args.aligns,
        "margin": args.margin,
        "valign": args.valign,
        "theme_style": args.theme_style,
        "clear_fill": args.clear_fill,
        "header_height": getattr(args, "header_height", None),
        "first_col_bold": getattr(args, "first_col_bold", None),
        "first_col_color": getattr(args, "first_col_color", None),
        "header_border": getattr(args, "header_border", None),
        "cell_runs": parse_json_arg(args.cell_runs, "--cell-runs") if getattr(args, "cell_runs", None) else None,
    }
    result = op_table_style(prs, spec)
    prs2, target = save_and_reload(prs, path, args.out)
    out = deck_meta(prs2, target)
    out.update(result)
    emit(out)


def cmd_apply(args):
    prs, path = load(args.pptx)
    ops = parse_json_arg(args.ops, "--ops")
    if not isinstance(ops, list):
        fail("--ops 必须是 JSON 数组")
    results = []
    errors = []
    for i, op in enumerate(ops):
        try:
            result = apply_one(prs, op)
        except KitError as exc:
            errors.append({"index": i, "op": (op or {}).get("op") if isinstance(op, dict) else None, "error": str(exc)})
            if args.stop_on_error:
                break
        except Exception as exc:  # unexpected backend failure: still report which op
            errors.append(
                {
                    "index": i,
                    "op": (op or {}).get("op") if isinstance(op, dict) else None,
                    "error": f"{type(exc).__name__}: {exc}",
                }
            )
            if args.stop_on_error:
                break
        else:
            results.append({"index": i, **result})

    touched = sorted({r["slide"] for r in results if isinstance(r.get("slide"), int)})
    out = {
        "dry_run": bool(args.dry_run),
        "stop_on_error": bool(args.stop_on_error),
        "op_count": len(ops),
        "applied": len(results),
        "changed_slides": touched,
        "ops": results,
        "errors": errors,
        "ok": not errors,
    }
    if errors and args.stop_on_error:
        out["saved"] = False
        out["hint"] = "批次已中止，文件未修改（原子提交）。修正 errors 后重试。"
        emit(out)
        return
    if args.dry_run:
        out["saved"] = False
        if args.shot and touched:
            try:
                ignore = {s.strip() for s in (args.ignore or "").split(",") if s.strip()}
                pages = [p for p in touched if 1 <= p <= len(prs.slides)]
                warnings = preview_overflow(prs, pages, args.overflow, ignore)
                if warnings:
                    out["overflow"] = warnings
                    out["overflow_preview"] = True
            except Exception as exc:  # rendering must never hide the validation result
                out["shot_error"] = str(exc)
        emit(out)
        return

    prs2, target = save_and_reload(prs, path, args.out)
    out.update(deck_meta(prs2, target))
    out["saved"] = True
    if args.shot and touched:
        try:
            pdf, _key, _cached = deck_pdf(target)
            ignore = {s.strip() for s in (args.ignore or "").split(",") if s.strip()}
            pages = [p for p in touched if 1 <= p <= len(prs2.slides)]
            warnings = render_overflow(prs2, pdf, pages, args.overflow, ignore)
            if warnings:
                out["overflow"] = warnings
        except Exception as exc:  # rendering must never lose the edit result
            out["shot_error"] = str(exc)
    emit(out)


def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description="PPTX agent toolkit")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("inspect", help="版式/形状 JSON")
    p.add_argument("pptx")
    p.add_argument("--slides", help="如 1,3 或 2-4；默认全部")
    p.add_argument("--mode", choices=["full", "summary"], default="full", help="summary 只给每页概览")
    p.add_argument("--depth", type=int, default=99, help="组合递归深度（0 = 不展开子形状）")
    p.set_defaults(func=cmd_inspect)

    p = sub.add_parser("shot", help="渲染页面为 PNG")
    p.add_argument("pptx")
    p.add_argument("-o", "--out", help="输出目录（默认放在缓存里）")
    p.add_argument("--slides", help="如 1,3 或 2-4；默认全部")
    p.add_argument("--dpi", type=int, default=110)
    p.add_argument("--max-pages", type=int, default=0, help="最多返回多少页")
    p.add_argument("--overflow", choices=["loose", "strict", "off"], default="loose",
                   help="溢出启发式：loose 跳过居中文本的左右误报（默认），strict 全部报告，off 关闭")
    p.add_argument("--ignore", help="跳过检查的形状：名称、#id 或 /regex/，逗号分隔")
    p.add_argument("--montage", type=int, metavar="COLS", help="额外输出拼图，如 3（3 列）")
    p.add_argument("--region", help="局部裁剪：x,y,w,h（英寸）或 JSON {\"3\": [x,y,w,h], \"default\": [...]}")
    p.add_argument("--only-changed", action=argparse.BooleanOptionalAction, default=False,
                   help="只渲染 --changed-slides 里列出的改动页（差分渲染）")
    p.add_argument("--changed-slides", help="改动页，如 3,7,9（配合 --only-changed）")
    p.add_argument("--density", action="store_true", help="附带每页首行与行数密度表")
    p.set_defaults(func=cmd_shot)

    p = sub.add_parser("create", help="新建/复制模板生成 pptx")
    p.add_argument("out")
    p.add_argument("--template")
    p.add_argument("--width", type=float, help="英寸")
    p.add_argument("--height", type=float, help="英寸")
    p.add_argument("--empty", action="store_true", help="不生成标题页")
    p.add_argument("--overwrite", action="store_true")
    p.set_defaults(func=cmd_create)

    p = sub.add_parser("clone", help="克隆已有页")
    p.add_argument("pptx")
    p.add_argument("src", type=int, help="源页号(1 起)")
    p.add_argument("--at", type=int, help="插入到第几页(1 起)")
    p.add_argument("--count", type=int, default=1)
    p.add_argument("--out")
    p.set_defaults(func=cmd_clone)

    p = sub.add_parser("add-slide", help="按版式新建页")
    p.add_argument("pptx")
    p.add_argument("--layout", required=True, help="版式序号或名称")
    p.add_argument("--out")
    p.set_defaults(func=cmd_add_slide)

    p = sub.add_parser("delete", help="删除页")
    p.add_argument("pptx")
    p.add_argument("--slides", required=True)
    p.add_argument("--out")
    p.set_defaults(func=cmd_delete)

    p = sub.add_parser("set-text", help="替换形状文本（保留原样式，可混排 runs）")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--shape", required=True, help="名称、路径、顶层序号或 #id")
    p.add_argument("--text", help="用 \\n 分段（与 --paras 二选一）")
    p.add_argument("--paras", help="段落/run 级样式 JSON 数组；@file 或 - 读文件/stdin")
    p.add_argument("--name", help="改名的同时给形状一个稳定句柄")
    p.add_argument("--out")
    p.set_defaults(func=cmd_set_text)

    p = sub.add_parser("set-cell", help="修改表格单元格")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--shape", required=True, help="名称、路径、顶层序号或 #id")
    p.add_argument("--row", type=int, required=True, help="0 起")
    p.add_argument("--col", type=int, required=True, help="0 起")
    p.add_argument("--text", required=True)
    p.add_argument("--out")
    p.set_defaults(func=cmd_set_cell)

    p = sub.add_parser("add-text", help="新增文本框")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--text", help="文本；\\n 分段")
    p.add_argument("--paras", help="段落/run 级样式 JSON 数组")
    p.add_argument("--x", type=float, required=True)
    p.add_argument("--y", type=float, required=True)
    p.add_argument("--w", type=float, default=6.0)
    p.add_argument("--h", type=float, help="省略时按内容估算并启用 grow")
    p.add_argument("--size", type=float)
    p.add_argument("--font")
    p.add_argument("--color")
    p.add_argument("--bold", action=argparse.BooleanOptionalAction, default=None)
    p.add_argument("--align", help="left|center|right|justify")
    p.add_argument("--valign", help="top|middle|bottom")
    p.add_argument("--line-spacing", type=float)
    p.add_argument("--space-after", type=float, help="段后间距 pt")
    p.add_argument("--space-before", type=float, help="段前间距 pt")
    p.add_argument("--indent", type=float, help="左缩进英寸；负值表示悬挂缩进")
    p.add_argument("--hanging", type=float, help="首行悬挂英寸（默认等于 indent）")
    p.add_argument("--bullet", help="项目符号，如 '•'；传 none 关闭")
    p.add_argument("--bullet-color")
    p.add_argument("--letter-spacing", type=float, help="字距（pt，写入 a:rPr/@spc）")
    p.add_argument("--name", help="文本框名（稳定的后续句柄）")
    p.add_argument("--fit", choices=["grow", "shrink", "none"], help="默认：省略 h 时 grow，否则 none")
    p.add_argument("--out")
    p.set_defaults(func=cmd_add_text)

    p = sub.add_parser("add-image", help="插入图片")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--path", required=True)
    p.add_argument("--x", type=float)
    p.add_argument("--y", type=float)
    p.add_argument("--w", type=float)
    p.add_argument("--h", type=float)
    p.add_argument("--out")
    p.set_defaults(func=cmd_add_image)

    p = sub.add_parser("add-table", help="插入表格")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--rows", type=int, required=True)
    p.add_argument("--cols", type=int, required=True)
    p.add_argument("--x", type=float, required=True)
    p.add_argument("--y", type=float, required=True)
    p.add_argument("--w", type=float, required=True)
    p.add_argument("--h", type=float, required=True)
    p.add_argument("--data", help="二维 JSON 数组；@file 亦可")
    p.add_argument("--size", type=float)
    p.add_argument("--header", action=argparse.BooleanOptionalAction, default=True)
    p.add_argument("--theme-style", action=argparse.BooleanOptionalAction, default=False,
                   help="沿用主题表格样式（默认关闭，避免与设计模板打架）")
    p.add_argument("--header-fill")
    p.add_argument("--header-color")
    p.add_argument("--zebra", help="偶数行填充色")
    p.add_argument("--border", help="边框颜色或 none")
    p.add_argument("--col-widths", help="逗号分隔的列宽（英寸）")
    p.add_argument("--row-height", type=float)
    p.add_argument("--header-height", type=float, help="表头行高（与 --row-height 独立）")
    p.add_argument("--first-col-bold", action=argparse.BooleanOptionalAction, default=None,
                   help="首列加粗（强调行标签）")
    p.add_argument("--first-col-color", help="首列字色")
    p.add_argument("--cell-runs", help="单元格富文本 JSON：[[[{t,bold,color}],...],...]")
    p.add_argument("--name", help="表格名（稳定的后续句柄）")
    p.add_argument("--aligns", help="逗号分隔的列对齐，如 left,center")
    p.add_argument("--out")
    p.set_defaults(func=cmd_add_table)

    p = sub.add_parser("add-shape", help="新增图形（圆角卡片/色条/箭头/分隔线）")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--kind", default="roundrect",
                   help="rect|roundrect|chevron|arrow|line|oval|triangle|hexagon…（默认 roundrect）")
    p.add_argument("--x", type=float, required=True)
    p.add_argument("--y", type=float, required=True)
    p.add_argument("--w", type=float, required=True)
    p.add_argument("--h", type=float, help="省略时：line/bar 为 0.06，其余 0.6")
    p.add_argument("--radius", type=float, help="圆角半径（英寸，仅 roundrect）")
    p.add_argument("--fill", help="填充色，none 表示无填充")
    p.add_argument("--line", help="描边色，none 表示无边框")
    p.add_argument("--line-width", type=float, default=1.0)
    p.add_argument("--shadow", action=argparse.BooleanOptionalAction, default=None,
                   help="是否继承主题阴影（默认关闭）")
    p.add_argument("--text")
    p.add_argument("--paras", help="段落/run 级样式 JSON")
    p.add_argument("--size", type=float)
    p.add_argument("--font")
    p.add_argument("--color")
    p.add_argument("--bold", action=argparse.BooleanOptionalAction, default=None)
    p.add_argument("--align", help="left|center|right|justify（默认 center）")
    p.add_argument("--valign", help="top|middle|bottom（默认 middle）")
    p.add_argument("--letter-spacing", type=float, help="字距（pt）")
    p.add_argument("--name", help="形状名（稳定的后续句柄）")
    p.add_argument("--out")
    p.set_defaults(func=cmd_add_shape)

    p = sub.add_parser("set-shape", help="移动/缩放/换色/改名/调叠放次序")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--shape", required=True, help="名称、路径、顶层序号或 #id")
    p.add_argument("--x", type=float, help="新的左边距（英寸）")
    p.add_argument("--y", type=float, help="新的上边距（英寸）")
    p.add_argument("--w", type=float, help="新的宽度（英寸）")
    p.add_argument("--h", type=float, help="新的高度（英寸）")
    p.add_argument("--fill", help="填充色，none 表示无填充")
    p.add_argument("--line", help="描边色，none 表示无边框")
    p.add_argument("--line-width", type=float)
    p.add_argument("--rotation", type=float, help="旋转角度（度）")
    p.add_argument("--z", choices=["front", "back", "forward", "backward"], help="叠放次序")
    p.add_argument("--name", help="改名")
    p.add_argument("--out")
    p.set_defaults(func=cmd_set_shape)

    p = sub.add_parser("delete-shape", help="删除形状")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--shape", required=True, help="名称、路径、顶层序号或 #id")
    p.add_argument("--out")
    p.set_defaults(func=cmd_delete_shape)

    p = sub.add_parser("set-notes", help="写讲者备注")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--text", required=True)
    p.add_argument("--out")
    p.set_defaults(func=cmd_set_notes)

    p = sub.add_parser("table-style", help="表格样式（填充/边框/列宽/对齐）")
    p.add_argument("pptx")
    p.add_argument("--slide", type=int, required=True)
    p.add_argument("--shape", required=True, help="表格形状名/路径/#id")
    p.add_argument("--header", action=argparse.BooleanOptionalAction, default=True)
    p.add_argument("--header-fill")
    p.add_argument("--header-color")
    p.add_argument("--body-color")
    p.add_argument("--zebra", help="偶数行填充色")
    p.add_argument("--border", help="边框颜色或 none")
    p.add_argument("--border-width", type=float, default=1.0)
    p.add_argument("--body-size", type=float)
    p.add_argument("--header-size", type=float)
    p.add_argument("--header-bold", action=argparse.BooleanOptionalAction, default=True)
    p.add_argument("--body-bold", action=argparse.BooleanOptionalAction, default=None)
    p.add_argument("--col-widths", help="逗号分隔的列宽（英寸）")
    p.add_argument("--row-height", type=float)
    p.add_argument("--header-height", type=float, help="表头行高（与 --row-height 独立）")
    p.add_argument("--first-col-bold", action=argparse.BooleanOptionalAction, default=None)
    p.add_argument("--first-col-color")
    p.add_argument("--header-border", help="表头单独边框色")
    p.add_argument("--cell-runs", help="单元格富文本 JSON")
    p.add_argument("--aligns", help="逗号分隔的列对齐，如 left,center")
    p.add_argument("--margin", type=float, help="单元格横向内边距（英寸）")
    p.add_argument("--valign", help="top|middle|bottom")
    p.add_argument("--theme-style", action=argparse.BooleanOptionalAction, default=None,
                   help="是否沿用主题表格样式（默认不改）")
    p.add_argument("--clear-fill", action="store_true", help="清除非表头行的填充")
    p.add_argument("--out")
    p.set_defaults(func=cmd_table_style)

    p = sub.add_parser("notes-check", help="逐页备注指纹（发现备注整体错位一页）")
    p.add_argument("pptx")
    p.add_argument("--slides")
    p.set_defaults(func=cmd_notes_check)

    p = sub.add_parser("apply", help="批量操作（一次事务提交 + 溢出汇总）")
    p.add_argument("pptx")
    p.add_argument("--ops", required=True, help="JSON 数组；@file 或 - 读文件/stdin")
    p.add_argument("--out")
    p.add_argument("--stop-on-error", action=argparse.BooleanOptionalAction, default=True)
    p.add_argument("--dry-run", action="store_true", help="只验证，不写文件")
    p.add_argument("--shot", action=argparse.BooleanOptionalAction, default=True,
                   help="保存后自动渲染改动页并汇总溢出/重叠")
    p.add_argument("--overflow", choices=["loose", "strict", "off"], default="loose")
    p.add_argument("--ignore", help="跳过检查的形状：名称、#id 或 /regex/，逗号分隔")
    p.set_defaults(func=cmd_apply)

    p = sub.add_parser("doctor", help="环境自检")
    p.add_argument("pptx", nargs="?", help="可选：同时读取该文件的主题字体")
    p.set_defaults(func=cmd_doctor)

    return ap


def main() -> None:
    ap = build_parser()
    args = ap.parse_args()
    try:
        args.func(args)
    except SystemExit:
        raise
    except KitError as exc:
        print(f"pptkit {args.cmd} 失败: {exc}", file=sys.stderr)
        raise SystemExit(2)
    except BrokenPipeError:
        pass
    except Exception as exc:
        if os.environ.get("PI_PPTKIT_DEBUG"):
            raise
        print(f"pptkit {args.cmd} 失败: {exc}", file=sys.stderr)
        raise SystemExit(2)


if __name__ == "__main__":
    main()
