"""Shape tree walking, shape resolution and per-slide inspection."""
from __future__ import annotations

import difflib

from pptx.enum.shapes import MSO_SHAPE_TYPE

from .utils import EMU, box_in, fail, parse_pages

def iter_shapes(shapes, prefix: str = ""):
    """Yield (shape, path). Group children are addressed as 'Group 5/TextBox 2'."""
    for shape in shapes:
        path = f"{prefix}{shape.name}"
        yield shape, path
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes, path + "/")


def _candidate_line(shape, path: str, index: int | None, tf=None) -> str:
    box = _abs_box(shape, tf or IDENTITY_TF)
    box_txt = f"box={box}" if box else "box=?"
    text = ""
    if shape.has_text_frame:
        text = shape.text_frame.text.strip().replace("\n", " / ")
        if len(text) > 24:
            text = text[:24] + "…"
    path_txt = f" path={path!r}" if path != shape.name else ""
    return (
        f"  index={index if index is not None else '-'} id=#{shape.shape_id} "
        f"name={shape.name!r}{path_txt} type={str(shape.shape_type).split(' ')[0]} {box_txt}"
        + (f" text={text!r}" if text else "")
    )


def _candidate_table(slide, limit: int = 20) -> str:
    rows = ["可用形状路径/候选形状（用 shape=\"#<id>\" 精确定位）:"]
    top = list(slide.shapes)
    for i, shape in enumerate(top[:limit]):
        rows.append(_candidate_line(shape, shape.name, i))
    if len(top) > limit:
        rows.append(f"  … 还有 {len(top) - limit} 个顶层形状")
    return "\n".join(rows)


def resolve_shape(slide, key: str):
    """Resolve a shape by '#id', exact path, unique name, or top-level index.

    Indexes are 0-based. Duplicate names are refused instead of silently
    picking the first match: templates often have several shapes named
    '标题 3'. Use the '#<id>' form returned by ppt_inspect in that case.
    """
    key = str(key).strip()
    top = list(slide.shapes)
    all_shapes = list(iter_shapes(slide.shapes))

    if key.startswith("#"):
        id_part = key[1:].strip()
        if not id_part.isdigit():
            fail(f"形状 id 必须是数字，如 '#12'，收到: {key}")
        wanted = int(id_part)
        matches = [(s, p) for s, p in all_shapes if s.shape_id == wanted]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            fail(f"找不到 id={wanted} 的形状。\n{_candidate_table(slide)}")
        fail(f"id={wanted} 在这页上不唯一，请检查 ppt_inspect 输出。")

    # exact recursive path
    exact = [(s, p) for s, p in all_shapes if p == key]
    if len(exact) == 1:
        return exact[0]

    # name match
    by_name = [(s, p) for s, p in all_shapes if s.name == key]
    if len(by_name) == 1:
        return by_name[0]
    if len(by_name) > 1:
        rows = [f"形状名 {key!r} 不唯一（{len(by_name)} 个），请改用 #<id>："]
        for shape, path in by_name:
            index = next(
                (i for i, s in enumerate(top) if s._element is shape._element),
                None,
            )
            rows.append(_candidate_line(shape, path, index))
        fail("\n".join(rows))

    # 0-based top-level index
    if key.lstrip("-").isdigit():
        idx = int(key)
        if -len(top) <= idx < len(top):
            shape = top[idx]
            return shape, shape.name
        fail(
            f"顶层序号 {idx} 超出范围 -{len(top)}..{len(top) - 1}（index 从 0 开始，本页共 {len(top)} 个顶层形状）。"
            f"提示：可改用 ppt_inspect 返回的 #id。\n{_candidate_table(slide)}"
        )

    # unique leaf name inside a group
    leaves = [(s, p) for s, p in all_shapes if s.name == key.rsplit("/", 1)[-1]]
    if len(leaves) == 1:
        return leaves[0]

    paths = [p for _, p in all_shapes]
    close = difflib.get_close_matches(key, paths, n=3, cutoff=0.35)
    hint = f"\n最相近的候选: {close}" if close else ""
    fail(f"找不到形状 {key!r}。{hint}\n{_candidate_table(slide)}")


def text_style(text_frame) -> dict | None:
    for para in text_frame.paragraphs:
        for run in para.runs:
            f = run.font
            style = {}
            if f.size is not None:
                style["size_pt"] = round(f.size.pt, 1)
            if f.name:
                style["font"] = f.name
            if f.bold is not None:
                style["bold"] = bool(f.bold)
            if f.italic is not None:
                style["italic"] = bool(f.italic)
            if para.alignment is not None:
                style["align"] = str(para.alignment).split(" ")[0].lower()
            return style or None
    return None


IDENTITY_TF = (1.0, 1.0, 0.0, 0.0)


def _abs_box(shape, tf) -> list[float] | None:
    sx, sy, ox, oy = tf
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return [
        round((shape.left * sx + ox) / EMU, 3),
        round((shape.top * sy + oy) / EMU, 3),
        round(shape.width * sx / EMU, 3),
        round(shape.height * sy / EMU, 3),
    ]


def _group_child_tf(group, tf):
    """Map a group's child coordinate space onto the slide canvas."""
    try:
        xfrm = group._element.grpSpPr.xfrm
        ch_off, ch_ext = xfrm.chOff, xfrm.chExt
        chex, chey = ch_ext.cx, ch_ext.cy
        if not chex or not chey:
            return tf
    except Exception:
        return tf
    sx, sy, ox, oy = tf
    gx, gy = group.left * sx + ox, group.top * sy + oy
    gw, gh = group.width * sx, group.height * sy
    csx, csy = gw / chex, gh / chey
    return (csx, csy, gx - ch_off.x * csx, gy - ch_off.y * csy)


def shape_info(shape, path: str, top_index: int | None = None, tf=IDENTITY_TF, depth: int = 99) -> dict:
    info = {
        "name": shape.name,
        "path": path,
        "type": str(shape.shape_type).split(" ")[0],
        "id": shape.shape_id,
    }
    if top_index is not None:
        # Usable as `shape` in the editing tools. Matters when a slide has
        # duplicate shape names (common when a template hand-places text boxes).
        info["index"] = top_index
    box = _abs_box(shape, tf)
    if box:
        info["box_in"] = box
    if shape.has_text_frame and shape.text_frame.text.strip():
        info["text"] = shape.text_frame.text
        style = text_style(shape.text_frame)
        if style:
            info["style"] = style
    if getattr(shape, "has_table", False) and shape.has_table:
        info["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
    if getattr(shape, "has_chart", False) and shape.has_chart:
        info["chart"] = True
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image = getattr(shape, "image", None)
        if image is not None:
            info["image"] = {"ext": image.ext, "px": [image.size[0], image.size[1]]}
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP and depth > 0:
        child_tf = _group_child_tf(shape, tf)
        info["children"] = [
            shape_info(child, f"{path}/{child.name}", None, child_tf, depth - 1) for child in shape.shapes
        ]
    return info


def shape_brief(shape, index: int | None = None) -> dict:
    info = {
        "id": shape.shape_id,
        "name": shape.name,
        "type": str(shape.shape_type).split(" ")[0],
        "box_in": box_in(shape),
    }
    if index is not None:
        info["index"] = index
    if shape.is_placeholder:
        info["placeholder_idx"] = shape.placeholder_format.idx
        info["placeholder_type"] = str(shape.placeholder_format.type).split(" ")[0]
    return info


def slide_title(slide) -> str:
    texts = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        name = shape.name.lower()
        if "标题" in name or "title" in name:
            return text.replace("\n", " ")[:60]
        texts.append(text.replace("\n", " "))
    return (texts[0][:60] if texts else "")


def slide_summary(prs, idx: int) -> dict:
    slide = prs.slides[idx]
    types: dict[str, int] = {}
    has_table = has_picture = has_chart = False
    for shape in slide.shapes:
        kind = str(shape.shape_type).split(" ")[0]
        types[kind] = types.get(kind, 0) + 1
        has_table = has_table or bool(getattr(shape, "has_table", False) and shape.has_table)
        has_chart = has_chart or bool(getattr(shape, "has_chart", False) and shape.has_chart)
        has_picture = has_picture or shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    info = {
        "index": idx + 1,
        "layout": slide.slide_layout.name,
        "title": slide_title(slide),
        "shape_count": len(slide.shapes),
        "types": types,
        "has_table": has_table,
        "has_picture": has_picture,
        "has_chart": has_chart,
    }
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            info["notes"] = notes[:40] + ("…" if len(notes) > 40 else "")
    return info


def slide_info(prs, idx: int, depth: int = 99) -> dict:
    slide = prs.slides[idx]
    info = {
        "index": idx + 1,
        "layout": slide.slide_layout.name,
        "layout_index": list(prs.slide_layouts).index(slide.slide_layout),
        "shape_count": len(slide.shapes),
        "shapes": [
            shape_info(shape, shape.name, i, depth=depth) for i, shape in enumerate(slide.shapes)
        ],
    }
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text
        if notes.strip():
            info["notes"] = notes
    return info


# ------------------------------------------------------------------ inspection
def layout_shapes(layout) -> list[dict]:
    """Non-placeholder shapes that the layout draws behind every slide.

    Templates keep the logo, rules and full-bleed background in the layout, so
    ``placeholders`` alone cannot answer "is this design inherited or per-slide".
    """
    out = []
    for shape in layout.shapes:
        if shape.is_placeholder:
            continue
        entry = {
            "name": shape.name,
            "type": str(shape.shape_type).split(" ")[0],
            "box_in": box_in(shape),
            "inherited": True,
        }
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = getattr(shape, "image", None)
            if image is not None:
                entry["image"] = {"ext": image.ext, "px": [image.size[0], image.size[1]]}
        out.append(entry)
    return out


def layout_catalog(prs) -> list[dict]:
    catalog = []
    for i, layout in enumerate(prs.slide_layouts):
        shapes = layout_shapes(layout)
        entry = {
            "index": i,
            "name": layout.name,
            "placeholders": [
                {
                    "idx": ph.placeholder_format.idx,
                    "type": str(ph.placeholder_format.type).split(" ")[0],
                    "name": ph.name,
                }
                for ph in layout.placeholders
            ],
        }
        if shapes:
            entry["shapes"] = shapes
            entry["shapes_hint"] = "继承自版式，勿在页内重复添加"
        catalog.append(entry)
    return catalog


def inspect_deck(prs, path, *, mode: str = "full", slides=None, depth: int = 99) -> dict:
    from .utils import deck_meta

    total = len(prs.slides)
    pages = parse_pages(slides, total) if slides else list(range(1, total + 1))
    depth = max(0, int(depth))
    out = deck_meta(prs, path)
    out["mode"] = mode
    out["layouts"] = layout_catalog(prs)
    if mode == "summary":
        out["slides"] = [slide_summary(prs, p - 1) for p in pages]
    else:
        out["slides"] = [slide_info(prs, p - 1, depth) for p in pages]
    out["inspected_slides"] = pages
    return out


def notes_check(prs, slides=None) -> dict:
    """Per-slide notes fingerprint: catches notes shifted by one page.

    Deleting a slide after writing notes is the classic silent breakage: the
    notes stay in order but every one of them now belongs to the previous page.
    """
    total = len(prs.slides)
    pages = parse_pages(slides, total) if slides else list(range(1, total + 1))
    rows = []
    for page in pages:
        slide = prs.slides[page - 1]
        text = ""
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
        title = slide_title(slide)
        rows.append(
            {
                "slide": page,
                "notes_len": len(text),
                "first_20": text[:20],
                "title": title[:20],
                "match": bool(text) and bool(title) and bool(set(text[:20]) & set(title)),
            }
        )
    empty = [r["slide"] for r in rows if r["notes_len"] == 0]
    out = {"slide_count": total, "notes": rows}
    if empty:
        out["slides_without_notes"] = empty
    return out
