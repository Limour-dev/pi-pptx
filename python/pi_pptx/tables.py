"""Table styling: theme style toggle, fills, borders, fonts, column widths."""
from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .utils import ALIGN, VALIGN, fail, parse_enum, parse_rgb

def _table_element(shape):
    return shape._element.graphic.graphicData.tbl


def _tbl_pr(tbl):
    tblPr = tbl.find(qn("a:tblPr"))
    if tblPr is None:
        tblPr = tbl.makeelement(qn("a:tblPr"), {})
        tbl.insert(0, tblPr)
    return tblPr


DEFAULT_TABLE_STYLE = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"


def set_theme_table_style(shape, enabled: bool) -> None:
    """Turn the theme's firstRow/bandRow variants and tableStyleId on/off."""
    tbl = _table_element(shape)
    tblPr = _tbl_pr(tbl)
    if enabled:
        tblPr.set("firstRow", "1")
        tblPr.set("bandRow", "1")
        if tblPr.find(qn("a:tableStyleId")) is None:
            el = tblPr.makeelement(qn("a:tableStyleId"), {})
            el.text = DEFAULT_TABLE_STYLE
            tblPr.insert(0, el)
        return
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for child in list(tblPr):
        if child.tag == qn("a:tableStyleId"):
            tblPr.remove(child)


def _set_cell_borders(cell, color: str | None, width_pt: float = 1.0) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for idx, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.makeelement(
            qn(tag),
            {"w": str(int(width_pt * 12700)), "cap": "flat", "cmpd": "sng", "algn": "ctr"},
        )
        if color:
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(parse_rgb(color))}))
            ln.append(fill)
            ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.insert(idx, ln)


def style_table(shape, spec: dict) -> dict:
    table = shape.table
    nrows, ncols = len(table.rows), len(table.columns)

    if spec.get("theme_style") is not None:
        set_theme_table_style(shape, bool(spec["theme_style"]))

    col_widths = spec.get("col_widths")
    if isinstance(col_widths, str):
        col_widths = [v.strip() for v in col_widths.split(",") if v.strip()]
    if col_widths:
        try:
            widths = [float(v) for v in col_widths]
        except (TypeError, ValueError):
            fail(f"col_widths 必须是数字数组（英寸），收到: {col_widths}")
        if len(widths) != ncols:
            fail(f"col_widths 需要 {ncols} 个值，收到 {len(widths)}")
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    row_height = spec.get("row_height")
    if row_height is not None:
        for row in table.rows:
            row.height = Inches(float(row_height))
    # A header taller than the body is the norm in designed decks and cannot be
    # expressed with a single rowHeight.
    header_height = spec.get("header_height", spec.get("headerHeight"))
    if header_height is not None and nrows:
        table.rows[0].height = Inches(float(header_height))

    body_size = spec.get("body_size", spec.get("size"))
    header_size = spec.get("header_size", body_size)
    aligns = spec.get("aligns")
    if isinstance(aligns, str):
        aligns = [a.strip() for a in aligns.split(",")]
    if aligns:
        aligns = [parse_enum(a, ALIGN, "对齐方式") for a in aligns]
    margin = spec.get("margin")
    valign = parse_enum(spec.get("valign"), VALIGN, "垂直对齐")
    border = spec.get("border")
    bold_header = spec.get("header_bold", True)
    header_fill = spec.get("header_fill")
    header_color = spec.get("header_color")
    body_color = spec.get("body_color")
    zebra = spec.get("zebra")
    first_col_bold = spec.get("first_col_bold", spec.get("firstColBold"))
    first_col_color = spec.get("first_col_color", spec.get("firstColColor"))
    header_border = spec.get("header_border", spec.get("headerBorder"))
    cell_runs = spec.get("cell_runs", spec.get("cellRuns"))
    if cell_runs is not None and not isinstance(cell_runs, list):
        fail("cell_runs 必须是二维数组：cell_runs[row][col] = [{t, bold, color, ...}]")

    for r in range(nrows):
        for c in range(ncols):
            cell = table.cell(r, c)
            is_header = r == 0 and bool(spec.get("header", True))
            if is_header and header_fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = parse_rgb(str(header_fill))
            elif not is_header and zebra and r % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = parse_rgb(str(zebra))
            elif spec.get("clear_fill"):
                cell.fill.background()
            edge_color = border
            if is_header and header_border is not None:
                edge_color = header_border
            if edge_color is not None:
                _set_cell_borders(
                    cell,
                    None if edge_color in (False, "none", "false") else str(edge_color),
                    float(spec.get("border_width", 1.0)),
                )
            if margin is not None:
                cell.margin_left = Inches(float(margin))
                cell.margin_right = Inches(float(margin))
                cell.margin_top = Inches(float(margin) / 2)
                cell.margin_bottom = Inches(float(margin) / 2)
            if valign is not None:
                cell.vertical_anchor = valign
            rich = None
            if cell_runs and r < nrows and c < ncols and isinstance(cell_runs[r], list):
                if c < len(cell_runs[r]):
                    rich = cell_runs[r][c]
            if rich is not None:
                write_cell_runs(cell, rich)
            emphasize = first_col_bold is not None and c == 0 and not is_header
            for para in cell.text_frame.paragraphs:
                if aligns:
                    para.alignment = aligns[c % len(aligns)]
                elif is_header:
                    para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    size = header_size if is_header else body_size
                    if size is not None:
                        run.font.size = Pt(float(size))
                    color = header_color if is_header else body_color
                    if color:
                        run.font.color.rgb = parse_rgb(str(color))
                    if is_header and bold_header:
                        run.font.bold = True
                    elif emphasize and first_col_bold:
                        run.font.bold = True
                    elif spec.get("body_bold") is not None:
                        run.font.bold = bool(spec["body_bold"])
                    if first_col_color and c == 0 and not is_header:
                        run.font.color.rgb = parse_rgb(str(first_col_color))
    result = {"rows": nrows, "cols": ncols}
    if cell_runs:
        result["rich_cells"] = True
    return result


def write_cell_runs(cell, runs: list) -> None:
    """Replace a cell's text with multiple styled runs (one paragraph)."""
    from .text import write_paragraphs

    if isinstance(runs, dict):
        runs = [runs]
    if not isinstance(runs, list):
        fail(f"cell_runs 的每一项必须是 run 数组或 run 对象，收到: {type(runs).__name__}")
    write_paragraphs(cell.text_frame, [{"runs": runs}])
